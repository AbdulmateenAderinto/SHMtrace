"""Anti-HLA-A2 antibody maturation (Devilder et al. 2019), clone sequences against the parental construct.

    python crm/devilder_clones.py        -> crm/out/devilder/

The study's Additional file 1 publishes the parental heavy and light variable regions and, for each selected
cluster of rounds 1 and 2, the nucleotide sequence together with the nucleotide mutations the authors
annotated. That annotation is the point of this analysis: it gives an independent answer to compare calls
against, on a defined reference, without any read-level processing.

This is a curated collection of *selected* clones, so it says nothing about population frequencies, the
unselected mutation distribution, replication or technical error, and none of those are computed. Each
published cluster is kept as its own sequence; nothing is collapsed to a consensus.

The parental sequence is the reference, which is what `ancestral_reference` is for: these clones were chosen
because they carry the mutations, so a change shared by most of them is the result, not germline.
"""
import json
import re
import sys
from pathlib import Path

CRM = Path(__file__).resolve().parent
REPO = CRM.parents[2]
sys.path.insert(0, str(REPO / "src"))
DATA = Path.home() / "Nanoshm/public_data/Devilder2019"
OUT = CRM / "out/devilder"
DNA = re.compile(r"^[ACGT]{60,}$")
MUT = re.compile(r"([ACGT])(\d+)([ACGT])")


def rows(table):
    for r in table.rows:
        yield [c.text.strip() for c in r.cells]


def extract():
    import docx
    d = docx.Document(str(DATA / "supp1.docx"))
    parental = {}
    for cells in rows(d.tables[1]):
        seq = next((c.upper() for c in cells if DNA.match(c.upper())), None)
        if seq:
            parental[cells[0].lower() or f"chain{len(parental)}"] = seq
    clones = {}
    for ti, rnd in ((2, "R1"), (4, "R2")):
        for cells in rows(d.tables[ti]):
            seq = next((c.upper() for c in cells if DNA.match(c.upper())), None)
            if not seq:
                continue
            name = cells[0] or cells[1]
            muts = next((c for c in cells if MUT.search(c) and not DNA.match(c.upper())), "")
            clones.setdefault(rnd, []).append({"cluster": name, "published_mutations": muts, "seq": seq})
    return parental, clones


def diffs(ref, seq):
    """Substitutions of `seq` against `ref`, 1-based, comparing the overlapping prefix only."""
    n = min(len(ref), len(seq))
    return [(i + 1, ref[i], seq[i]) for i in range(n) if ref[i] != seq[i] and seq[i] in "ACGT"]


def main():
    if not (DATA / "supp1.docx").exists():
        sys.exit(f"missing {DATA / 'supp1.docx'}")
    OUT.mkdir(parents=True, exist_ok=True)
    parental, clones = extract()
    heavy = parental.get("heavy")
    if not heavy:
        sys.exit(f"no parental heavy chain found; parsed chains: {list(parental)}")
    (OUT / "parental_heavy.fasta").write_text(">parental_heavy\n"
                                              + "\n".join(heavy[i:i + 70] for i in range(0, len(heavy), 70))
                                              + "\n")
    out = {"source": "Devilder et al. 2019 BMC Biotechnol, Additional file 1 (10.1186/s12896-019-0504-z)",
           "parental_heavy_bp": len(heavy), "rounds": {}}
    with open(OUT / "clones.fasta", "w") as fh:
        for rnd, items in clones.items():
            recovered = []
            for c in items:
                d = diffs(heavy, c["seq"])
                published = {f"{a}{p}{b}" for a, p, b in MUT.findall(c["published_mutations"])}
                found = {f"{r}{p}{s}" for p, r, s in d}
                recovered.append({"cluster": c["cluster"], "length": len(c["seq"]),
                                  "published_mutations": sorted(published),
                                  "recovered_substitutions": sorted(found),
                                  "published_not_recovered": sorted(published - found),
                                  "recovered_not_published": sorted(found - published)})
                fh.write(f">{rnd}_{c['cluster'] or 'unnamed'}\n"
                         + "\n".join(c["seq"][i:i + 70] for i in range(0, len(c["seq"]), 70)) + "\n")
            out["rounds"][rnd] = recovered
    # and through the product's own clone path, which is what the policy has to work in
    from shmtrace import project
    runs = {}
    for policy in (True, False):
        cfg = {"name": "devilder", "region": "sanger_clones", "platform": "sanger",
               "reference": str(OUT / "parental_heavy.fasta"),
               "samples": [{"name": "selected_clones", "path": str(OUT / "clones.fasta"),
                            "group": "selected", "control": False}],
               "overrides": {"end_trim_bp": 0, "exclude_indels": True,
                             "ancestral_reference": policy, "jobs": 1}}
        sm = project.public(project.run_project(cfg))["samples"][0]
        runs["on" if policy else "off"] = {"clones": sm["reads"], "unique_mutations": sm["unique_mutations"],
                                           "removed_by_frequency_rule": sm.get("high_fraction_sites_removed")}
    out["shmtrace_sanger_path"] = runs
    json.dump(out, open(OUT / "devilder_clones.json", "w"), indent=1)
    print(f"parental heavy chain: {len(heavy)} bp")
    for rnd, items in out["rounds"].items():
        with_ann = [c for c in items if c["published_mutations"]]
        exact = [c for c in with_ann if not c["published_not_recovered"] and not c["recovered_not_published"]]
        miss = sum(len(c["published_not_recovered"]) for c in with_ann)
        extra = sum(len(c["recovered_not_published"]) for c in with_ann)
        print(f"{rnd}: {len(items)} clusters, {len(with_ann)} with published mutations; "
              f"{len(exact)} match exactly; {miss} published not recovered; {extra} recovered not published")
    r = out["shmtrace_sanger_path"]
    print(f"through shmtrace's clone path: {r['on']['unique_mutations']} unique with the policy, "
          f"{r['off']['unique_mutations']} without "
          f"({r['off']['removed_by_frequency_rule']} removed by the frequency rule)")


if __name__ == "__main__":
    main()
