"""Table S2 (Excel): per-library calls and densities for every caller and dataset, and spike-in summaries.
    python docs/calibration/crm/make_tables.py  -> out/Table_S2_benchmark.xlsx"""
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).parent / "out"
LABEL = {"recommended": "SHMtrace", "geneious": "Geneious rule", "lofreq": "LoFreq (defaults)",
         "lofreq_nobaq": "LoFreq (-B, BAQ off)", "lofreq_permissive": "LoFreq (-B, p<=0.05 uncorrected, filters off)",
         "v0.1": "per-position null", "pbinom": "quality model only",
         "clairs": "ClairS (defaults, ont_r10_dorado_sup_5khz)", "clairs_permissive": "ClairS (snv_min_af 0.005)",
         "clair3": "Clair3 (defaults, r1041_e82_400bps_sup_v520)", "clair3_permissive": "Clair3 (snp_min_af 0.01)",
         "deepvariant": "DeepVariant (defaults, ONT_R104)",
         "deepvariant_permissive": "DeepVariant (vsc_min_fraction_snps 0.01, vsc_min_count_snps 2)",
         "deepsomatic": "DeepSomatic (defaults, ONT)",
         "deepsomatic_permissive": "DeepSomatic (vsc_min_fraction_snps 0.005, vsc_min_count_snps 2)"}
SETS = {"jh4": "J_H4 intron, this study", "n114": "N114 V region, this study", "su_lps": "Smu, LPS (PRJNA1481264)",
        "su_ch12": "Smu, CH12 (PRJNA1481264)", "pim1": "Pim1, LPS (PRJNA1481264), 2,000 reads",
        "pim1_deep": "Pim1, LPS (PRJNA1481264), all usable reads"}
wb = Workbook()
ws = wb.active
ws.title = "Real libraries"
ws.append(["Dataset", "Library", "AID-null control", "Caller", "Reads analysed", "Bases in window",
           "Distinct calls", "Density per 1e5 bases"])
for c in ws[1]:
    c.font = Font(bold=True)
for key, name in SETS.items():
    f = OUT / f"benchmark_real_{key}.json"
    if not f.exists():
        continue
    for r in json.load(open(f))["rows"]:
        ws.append([name, r["sample"], "yes" if r["control"] else "no", LABEL[r["mode"]], r["reads"], r["bases"],
                   r["calls"], round(r["density_x1e5"], 3) if r["density_x1e5"] is not None else None])
ws2 = wb.create_sheet("Spike-ins")
ks = [1, 2, 3, 5, 10, 20]
ws2.append(["Dataset", "Caller", "Libraries x replicates"] + [f"Recall, {k} read(s)" for k in ks]
           + ["Recall, all", "False calls per library", "Observed FDR"])
for c in ws2[1]:
    c.font = Font(bold=True)
for key in ("jh4", "n114", "su_lps", "pim1"):
    f = OUT / f"spikein_{key}.json"
    if not f.exists():
        continue
    rows = json.load(open(f))["rows"]
    for tag in ("_clairs", "_clair3", "_deepvariant", "_deepsomatic"):   # other callers, same planted mutations
        if (OUT / f"spikein_{key}{tag}.json").exists():
            rows += [r for r in json.load(open(OUT / f"spikein_{key}{tag}.json"))["rows"] if r["caller"] != "recommended"]
    for caller in ("recommended", "geneious", "lofreq", "lofreq_nobaq", "lofreq_permissive",
                   "clairs", "clairs_permissive", "clair3", "clair3_permissive", "deepvariant",
                   "deepvariant_permissive", "deepsomatic", "deepsomatic_permissive"):
        rs = [r for r in rows if r["caller"] == caller]
        if not rs:                               # a caller not run on this dataset
            continue
        rec = defaultdict(lambda: [0, 0])
        for r in rs:
            for k, (got, n) in r["recall_by_reads"].items():
                rec[int(k)][0] += got
                rec[int(k)][1] += n
        tp, fp = sum(r["true_calls"] for r in rs), sum(r["false_calls"] for r in rs)
        tot = sum(v[0] for v in rec.values()) / max(1, sum(v[1] for v in rec.values()))
        ws2.append([SETS.get(key, key).split(",")[0] if key != "pim1" else "Pim1 (PRJNA1481264)", LABEL[caller], len(rs)]
                   + [round(rec[k][0] / rec[k][1], 3) if rec[k][1] else None for k in ks]
                   + [round(tot, 3), round(fp / len(rs), 2), round(fp / (tp + fp), 3) if tp + fp else 0])
for w in (ws, ws2):
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))
wb.save(OUT / "Table_S2_benchmark.xlsx")
print(OUT / "Table_S2_benchmark.xlsx", ws.max_row - 1, "library rows")
