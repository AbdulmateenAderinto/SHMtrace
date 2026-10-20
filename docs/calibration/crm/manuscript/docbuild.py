"""A small document builder for the Cell Press manuscripts (Cell Reports Methods, STAR Protocols).

Text is written with light markup:
    *italic*   **bold**   ^{superscript}   _{subscript}   [@key] or [@key1,@key2]  (citations)
    \\*  a literal asterisk (significance stars in a legend), which is not markup
Citations become numbered superscripts in order of first appearance (Cell style); the reference
list is built from refs.json (resolved against PubMed by resolve_refs.py) plus MANUAL entries for
the few works PubMed does not index. Blocks are (kind, payload) tuples; see add_block.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).parent
REFS = json.load(open(HERE / "refs.json"))
MANUAL = {
    "benjamini1995": "Benjamini, Y., and Hochberg, Y. (1995). Controlling the false discovery rate: a practical and "
                     "powerful approach to multiple testing. J. R. Stat. Soc. Series B Stat. Methodol. 57, 289–300. "
                     "10.1111/j.2517-6161.1995.tb02031.x.",
    "elias2007": "Elias, J.E., and Gygi, S.P. (2007). Target-decoy search strategy for increased confidence in "
                 "large-scale protein identifications by mass spectrometry. Nat. Methods 4, 207–214. 10.1038/nmeth1019.",
    "hubert1985": "Hubert, L., and Arabie, P. (1985). Comparing partitions. J. Classif. 2, 193–218. 10.1007/BF01908075.",
    # Clair3: confirmed in the HKU repository (hub.hku.hk/handle/10722/324284)
    "zheng2022": "Zheng, Z., Li, S., Su, J., Leung, A.W.-S., Lam, T.-W., and Luo, R. (2022). Symphonizing pileup and "
                 "full-alignment for deep learning-based long-read variant calling. Nat. Comput. Sci. 2, 797–803. "
                 "10.1038/s43588-022-00387-x.",
    # ClairS: the Nature Methods paper (authors, volume and pages from Crossref; the bioRxiv preprint was
    # 10.1101/2023.08.17.553778)
    "zheng2026": "Zheng, Z., Chen, L., Su, J., Yu, X., He, M., Lee, Y.-L., Lam, T.-W., and Luo, R. (2026). ClairS: a "
                 "deep-learning method for long-read tumor–normal pair somatic small variant calling. Nat. Methods "
                 "23, 1507–1520. 10.1038/s41592-026-03152-4.",
    # Seija et al.: MLLT1/MLLT3 concentrate AID (Crossref; Cell Press style, first 10 authors then et al.)
    "seija2026": "Seija, N., Gannon, S., Häfner, K., Gemeinhardt, T., Ridani, J., Alvarez, D., Provencher, M., "
                 "Subramani, P., Poitras, C., Piskor, E., et al. (2026). Histone readers MLLT1 and MLLT3 "
                 "concentrate AID to confer locus specificity. Nature. 10.1038/s41586-026-11087-1.",
    # Perez-Duran et al.: ectopic AID in non-B cells (authors, volume and pages from Crossref)
    "perezduran2012": "Pérez-Durán, P., Belver, L., de Yébenes, V.G., Delgado, P., Pisano, D.G., and Ramiro, "
                      "A.R. (2012). UNG shapes the specificity of AID-induced somatic hypermutation. J. Exp. "
                      "Med. 209, 1379–1389. 10.1084/jem.20112253.",
    # DeepVariant and DeepSomatic: authors, volumes and pages from Crossref (Cell Press: first 10 authors, et al.)
    "poplin2018": "Poplin, R., Chang, P.-C., Alexander, D., Schwartz, S., Colthurst, T., Ku, A., Newburger, D., "
                  "Dijamco, J., Nguyen, N., Afshar, P.T., et al. (2018). A universal SNP and small-indel variant "
                  "caller using deep neural networks. Nat. Biotechnol. 36, 983–987. 10.1038/nbt.4235.",
    "park2026": "Park, J., Cook, D.E., Chang, P.-C., Kolesnikov, A., Brambrink, L., Mier, J.C., Gardner, J., "
                "McNulty, B., Sacco, S., Keskus, A.G., et al. (2026). Accurate somatic small variant discovery for "
                "multiple sequencing technologies with DeepSomatic. Nat. Biotechnol. 44, 1569–1578. "
                "10.1038/s41587-025-02839-x.",
}
PAGE_FIX = {"sanderson2024": "001246"}
LITERAL_STAR = "\uE000"     # stands in for an escaped asterisk while the markup is read
TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\^\{[^}]*\}|_\{[^}]*\}|\[@[^\]]+\])")
# figure and table citations in running text, for the linked build: "Figure 3", "Figures 4A–4E and 4G",
# "Figures S8–S10", "Tables S4 and S5". Each item links to its figure or table; a noun links with its first item.
_ITEM = r"S?\d+[A-Z]*(?:–S?\d+[A-Z]*)?"
CITATION = re.compile(rf"\b(Figures?|Tables?)\s({_ITEM}(?:(?:,\s+and\s+|,\s+|\s+and\s+)(?:{_ITEM}))*)")
TITLE_ANCHOR = re.compile(r"^\**(Figure|Table) (S?\d+)\.")
# panel labels in a legend, set bold inside plain parentheses as the preprint does: (A), (B, C), (A–E), (D and E)
PANELS = re.compile(r"\(([A-H](?:–[A-H])?(?:(?:,\s*|\s+and\s+)[A-H](?:–[A-H])?)*)\)")
LINK_RGB = RGBColor(0x1F, 0x5F, 0xA8)


def anchor(kind, number):
    """Bookmark name for a figure, table or reference: fig_3, fig_S4, tab_S2, ref_12."""
    return f"{'fig' if kind.startswith('Fig') else 'tab' if kind.startswith('Tab') else 'ref'}_{number}"


class Manuscript:
    def __init__(self, font="Arial", size=11, line_numbers=True, links=False, footer_note=None):
        # links: figure, table and reference citations become internal hyperlinks to bookmarks on the
        # figure and table titles and the reference entries (the combined review document)
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = font, Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
        pf = st.paragraph_format
        pf.line_spacing, pf.space_after = 1.5, Pt(6)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Mm(279.4), Mm(215.9)      # US letter (Cell Press, US usage)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Mm(25.4))
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            sec._sectPr.append(ln)
        self.links, self._bookmarks, self._next_id = links, set(), 1
        self._page_numbers(sec, footer_note)
        self._page_break_next = False
        self.order = []          # citation keys in order of first appearance
        self.font = font

    def _page_numbers(self, sec, note=None):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if note:
            p.add_run(f"{note} · page ").font.size = Pt(9)
        r = p.add_run()
        for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if t:
                f = OxmlElement("w:fldChar")
                f.set(qn("w:fldCharType"), t)
                r._r.append(f)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._r.append(it)

    # ---- bookmarks and internal links
    def bookmark(self, par, name):
        """Mark a paragraph as the target `name`; the first paragraph to claim a name keeps it."""
        if not self.links or name in self._bookmarks:
            return
        self._bookmarks.add(name)
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(self._next_id))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(self._next_id))
        self._next_id += 1
        par._p.insert(1 if par._p.pPr is not None else 0, start)
        par._p.append(end)

    def _link(self, par, target):
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), target)
        h.set(qn("w:history"), "1")
        par._p.append(h)
        return h

    def _unlink_missing(self):
        """A link whose target is not in this document (the Excel table, Data S1) becomes plain text again."""
        for h in list(self.doc.element.body.iter(qn("w:hyperlink"))):
            if h.get(qn("w:anchor")) not in self._bookmarks:
                parent = h.getparent()
                for r in list(h):
                    r_pr = r.find(qn("w:rPr"))
                    color = r_pr.find(qn("w:color")) if r_pr is not None else None
                    if color is not None:
                        r_pr.remove(color)
                    parent.insert(parent.index(h), r)
                parent.remove(h)

    # ---- inline markup
    def cite(self, keys):
        nums = []
        for k in keys:
            if k not in REFS and k not in MANUAL:
                raise KeyError(f"unknown reference {k}")
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        nums = sorted(set(nums))
        out, i = [], 0
        while i < len(nums):              # 1,2,3 -> 1–3
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(str(x) for x in nums[i:j + 1]))
            i = j + 1
        return ",".join(out)

    def runs(self, par, text, size=None, bold=False, color=None, legend=False):
        # an escaped asterisk is hidden from the markup and printed as itself
        text = text.replace("\\*", LITERAL_STAR)
        for piece in TOKEN.split(text):
            if not piece:
                continue
            kw, cite_keys = {}, None
            if piece.startswith("`"):
                r = par.add_run(piece[1:-1].replace(LITERAL_STAR, "*"))
                r.font.name = "Courier New"
                r.font.size = Pt((size or 11) - 1)
                continue
            if piece.startswith("**"):
                piece, kw = piece[2:-2], {"bold": True}
            elif piece.startswith("*"):
                piece, kw = piece[1:-1], {"italic": True}
            elif piece.startswith("^{"):
                piece, kw = piece[2:-1], {"sup": True}
            elif piece.startswith("_{"):
                piece, kw = piece[2:-1], {"sub": True}
            elif piece.startswith("[@"):
                cite_keys = [k.strip().lstrip("@") for k in piece[2:-1].split(",")]
                piece, kw = self.cite(cite_keys), {"sup": True}
            # italics inside sup/sub etc. are not nested; keep markup flat
            sub_parts = TOKEN.split(piece) if kw and ("sup" in kw or "sub" in kw) and "*" in piece else [piece]
            for sp in sub_parts:
                if not sp:
                    continue
                txt = (sp.strip("*") if sp.startswith("*") else sp).replace(LITERAL_STAR, "*")
                italic = kw.get("italic") or sp.startswith("*")
                if self.links and cite_keys is not None:
                    # each reference number links to its entry: "3–5,8" -> 3, 5 and 8
                    for tok in re.split(r"(\d+)", txt):
                        if tok:
                            self._emit(par, tok, bold or kw.get("bold"), italic, kw, size, color,
                                       anchor("ref", tok) if tok.isdigit() else None)
                elif (self.links or legend) and not kw:
                    # figure and table citations (links in the linked build, bold in a legend); panel labels
                    # bold in a legend
                    spans = [(m.start(), m.end(), "cite", m) for m in CITATION.finditer(txt)]
                    if legend:
                        spans += [(m.start(), m.end(), "panel", m) for m in PANELS.finditer(txt)
                                  if not any(a <= m.start() < b for a, b, *_ in spans)]
                    pos = 0
                    for start, end, what, m in sorted(spans, key=lambda x: x[0]):
                        self._emit(par, txt[pos:start], bold, italic, kw, size, color)
                        if what == "panel":
                            self._emit(par, "(", bold, italic, kw, size, color)
                            self._emit(par, m.group(1), True, italic, kw, size, color)
                            self._emit(par, ")", bold, italic, kw, size, color)
                        else:
                            noun, items = m.group(1), m.group(2)
                            parts = re.split(r"(,\s+and\s+|,\s+|\s+and\s+|–)", items)
                            for k, tok in enumerate(parts):
                                num = re.match(r"S?\d+", tok)
                                target = anchor(noun, num.group(0)) if num and self.links else None
                                self._emit(par, (f"{noun} " if k == 0 else "") + tok, bold or legend, italic, kw,
                                           size, color, target)
                        pos = end
                    self._emit(par, txt[pos:], bold, italic, kw, size, color)
                else:
                    self._emit(par, txt, bold or kw.get("bold"), italic, kw, size, color)
        return par

    def _emit(self, par, txt, bold, italic, kw, size, color, target=None):
        if not txt:
            return None
        r = par.add_run(txt)
        r.bold = bold
        r.italic = italic
        # superscript and subscript share w:vertAlign: set only the one that applies
        if kw.get("sup"):
            r.font.superscript = True
        elif kw.get("sub"):
            r.font.subscript = True
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if target:
            r.font.color.rgb = LINK_RGB
            r._r.getparent().remove(r._r)
            self._link(par, target).append(r._r)
        return r

    # ---- blocks
    def _paragraph(self, *args, **kwargs):
        p = self.doc.add_paragraph(*args, **kwargs)
        if self._page_break_next:
            p.paragraph_format.page_break_before = True
            self._page_break_next = False
        return p

    def heading(self, text, level=1):
        sizes = {0: 16, 1: 13, 2: 11.5, 3: 11}
        p = self._paragraph()
        p.paragraph_format.space_before = Pt(12 if level <= 1 else 8)
        p.paragraph_format.keep_with_next = True
        self.runs(p, text, size=sizes[level], bold=True)
        return p

    def para(self, text, size=None, align=None, bold=False, indent=None, color=None, space_after=None):
        p = self._paragraph()
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if indent:
            p.paragraph_format.left_indent = Mm(indent)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        title = TITLE_ANCHOR.match(text)
        if title:
            self.bookmark(p, anchor(title.group(1), title.group(2)))
        return self.runs(p, text, size=size, bold=bold, color=color)

    def bullets(self, items, numbered=False, size=None):
        for it in items:
            p = self._paragraph(style="List Number" if numbered else "List Bullet")
            self.runs(p, it, size=size)

    def figure(self, path, width_mm=170):
        p = self._paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Mm(width_mm))
        return p

    def caption(self, title, legend, size=10):
        """A figure caption as the authors' preprint sets it: the bold title, ending in a full stop, runs into the
        legend in one justified paragraph; panel labels and figure or table references are bold."""
        title = title.rstrip()
        title = title if title.endswith((".", "?", "!")) else title + "."
        p = self._paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)
        found = TITLE_ANCHOR.match(title)
        if found:
            self.bookmark(p, anchor(found.group(1), found.group(2)))
        self.runs(p, title + " ", size=size, bold=True)
        return self.runs(p, legend, size=size, legend=True)

    def figure_block(self, paths, title, legend, max_width_mm=165, min_width_mm=125):
        """One figure and its complete legend on one page.

        A tall figure placed at full width pushes its legend onto a page of its own, which is both wasteful
        and hard to read. The images are scaled down, together, until the figure and its legend fit the text
        area, down to `min_width_mm`. Below that a figure drawn for 165 mm would carry text under the 6 pt
        floor for print, so legibility wins and the legend is allowed to run on instead.
        """
        from PIL import Image as _Image
        _Image.MAX_IMAGE_PIXELS = None
        sec = self.doc.sections[-1]
        usable = (sec.page_height - sec.top_margin - sec.bottom_margin) / 36000 - 12
        # Actual Letter text height, with room for image paragraphs and the page-break line.
        ratio = sum(_Image.open(str(p)).size[1] / _Image.open(str(p)).size[0] for p in paths)
        lines = max(1, round(len(str(legend).split()) / 11)) + 2      # legend plus its bold title
        legend_mm = lines * 4.6
        width = max_width_mm
        if ratio:
            width = min(max_width_mm, max(min_width_mm, (usable - legend_mm) / ratio))
        for p in paths:
            self.figure(p, width_mm=width)
        self.caption(title, legend)
        return width

    def table(self, header, rows, widths=None, size=8.5, title=None, footnote=None):
        if title:
            self.para(title, bold=False, size=10)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        header_props = t.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_props.append(repeat_header)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, h in zip(t.rows[0].cells, header, strict=True):
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            self.runs(c.paragraphs[0], h, size=size, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, row, strict=True):
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                c.paragraphs[0].paragraph_format.line_spacing = 1.0
                self.runs(c.paragraphs[0], str(v), size=size)
        if widths:
            for row in t.rows:
                for c, w in zip(row.cells, widths, strict=True):
                    c.width = Mm(w)
        if footnote:
            note = self.para(footnote, size=8.5, space_after=4)
            note.paragraph_format.line_spacing = 1.0
            note.paragraph_format.keep_together = True
        return t

    def page_break(self):
        self._page_break_next = True

    def references(self, heading="References"):
        self.heading(heading, 1)
        for i, k in enumerate(self.order, 1):
            if k in MANUAL:
                txt = MANUAL[k]
            else:
                r = REFS[k]
                txt = r["cell"]
                if k in PAGE_FIX:
                    txt = txt.replace(", . ", f", {PAGE_FIX[k]}. ")
            p = self._paragraph()
            p.paragraph_format.left_indent = Mm(7)
            p.paragraph_format.first_line_indent = Mm(-7)
            p.paragraph_format.line_spacing = 1.15
            self.bookmark(p, anchor("ref", i))
            self.runs(p, f"{i}.\t{txt}", size=10)

    def save(self, path):
        if self.links:
            self._unlink_missing()
        self.doc.save(path)


def word_count(blocks):
    n = 0
    for _k, t in blocks:
        if isinstance(t, str):
            n += len(TOKEN.sub(lambda m: "" if m.group().startswith("[@") else m.group().strip("*^_{}"), t).split())
    return n
