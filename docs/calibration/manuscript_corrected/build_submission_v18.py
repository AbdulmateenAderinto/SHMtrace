"""Assemble the Ver18 submission package and the Ver10 data deposit beside the earlier ones.

    python build_submission_v18.py          (needs python-docx and Pillow; LibreOffice for PDFs)

Nothing earlier is modified: the Ver13 package and Ver9 deposit are read and copied.
  ONT_SHM_JIM_Submission_Ver18_<date>/   manuscript, cover letter, highlights, figures,
                                         graphical abstract, code, data, supporting notes
  ONT_SHM_DataDeposit_Ver10_<date>/      Ver9 plus the Sanger chromatograms, the sample
                                         sheet and every SHMtrace analysis behind the figures
Run revise_manuscript_v18.py first.
"""
import csv
import hashlib
import shutil
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from PIL import Image

HERE = Path(__file__).resolve().parent
CAL = HERE.parent
REPO = CAL.parents[1]
FIG = HERE / "figures"
SV = CAL / "figures/sanger_vs_nanopore"
SUB = Path.home() / ("Documents/Projects/2026/PhD - Fundamental Immunology/Martin Lab/"
                     "Paper #1 - Nanopore sequencing/June 2nd analyses/submission_v4")
V13 = SUB / "ONT_SHM_JIM_Submission_Ver13_2026-08-01"
DEP9 = SUB / "ONT_SHM_DataDeposit_Ver9_2026-07-31"
MS = SUB / "ONT_SHM_Manuscript_JIM_Ver18_SHMtrace.docx"
AB1 = Path.home() / "Nanoshm/sanger_ab1_original"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
STAMP = date.today().isoformat()
PKG = SUB / f"ONT_SHM_JIM_Submission_Ver18_{STAMP}"
DEP = SUB / f"ONT_SHM_DataDeposit_Ver10_{STAMP}"
TITLE = ("Cloning-Free Profiling of Somatic Hypermutation with Oxford Nanopore R10.4.1/Kit 14 Sequencing: "
         "Proof of Concept in the Mouse JH4 Intron and an Immunoglobulin V Region")

# manuscript figure -> (PNG source, vector PDF source)
FIGURES = {
    "Figure1": FIG / "Figure1_study_design_pipeline",
    "Figure2": FIG / "Figure2_JH4_combined",
    "Figure3": FIG / "Figure3_concordance",
    "Figure4": FIG / "Figure4_hybridoma_combined",
    "FigureS1": FIG / "FigureS_spectrum_529_vs_fulllength",
    "FigureS2": FIG / "FigureS5_individual_Sanger_matrices",
    "FigureS3": FIG / "FigureS6_individual_JH4_ONT_matrices",
    "FigureS4": FIG / "FigureS7_individual_N114_matrices",
    "FigureS5": FIG / "FigureS7_recurrence_weighted",
    "FigureS6": FIG / "FigureS3_depth_discovery_density",
    "FigureS7": FIG / "FigureS4_saturation",
    "FigureS8": SV / "figS_sanger_window_original",
    "FigureS9": SV / "recommended_manuscript/fig3_agreement_position_hotspot",
}
# SHMtrace analyses deposited: run folder -> deposit name
RUNS = {"jh4_shmtrace": "JH4_529bp_window", "jh4_shmtrace_full": "JH4_full_length_90-1230",
        "jh4_shmtrace_raw": "JH4_delivered_reads_sensitivity", "n114_shmtrace": "N114_V_region",
        "depth_shmtrace": "Depth_series", "sanger_vs_nanopore/recommended_manuscript": "Sanger_vs_ONT"}
SANGER_MICE = ["WT2", "WT3", "WT4", "Aicda_A1", "Aicda_A2", "Aicda_A3"]

COVER = {
    0: f"{date.today():%B} {date.today().day}th, {date.today().year}",
    2: f"Please consider our manuscript by Abdulmateen Aderinto & Alberto Martin, entitled \"{TITLE},\" for "
       "publication in the Journal of Immunological Methods.",
    5: "We therefore developed a cloning-free nanopore workflow for SHM, implemented it as SHMtrace, an "
       "open-source analysis tool that tests every candidate substitution against matched negative-control "
       "libraries with false-discovery control, and benchmarked it directly against cloned-amplicon Sanger "
       "sequencing in the same animals, re-analysed from the original chromatograms.",
    6: "The two platforms agreed. Both recovered the expected AID-dependent difference between wild-type and "
       "Aicda−/− mice (at least 108-fold by nanopore, ~38-fold by Sanger), and their pooled mutation spectra "
       "correlated closely across the twelve substitution types (Pearson r = 0.91). An independent monoclonal "
       "system, an AID-expressing hybridoma line, gave at least an 89-fold difference over empty-vector "
       "controls with the expected G:C-biased signature, indicating that the framework is not specific to a "
       "single locus. Residual background in AID-deficient controls was ~0.5 × 10−5 per analysed base.",
    7: "Nanopore also yields information the cloning-based assay cannot. Because each read spans the entire "
       "amplicon, we could analyse ~1,140 bp rather than the ~530 bp window accessible to Sanger sequencing, "
       "and retain co-occurring mutations linked on individual reads.",
}
HIGHLIGHTS = {
    1: "Control-calibrated ONT calling profiles full-length SHM amplicons without cloning",
    2: "Matched biological controls define a Q53-equivalent residual background",
    3: "ONT reproduces the paired Sanger JH4 substitution spectrum (r = 0.91)",
    4: "Approximately 500 reads recover ~91% of discoverable AID-active variants",
}


def set_paragraph(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def to_pdf(docx, outdir):
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", str(docx), "--outdir", str(outdir)],
                   check=True, capture_output=True)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def build_deposit():
    if DEP.exists():
        shutil.rmtree(DEP)
    shutil.copytree(DEP9, DEP, ignore=shutil.ignore_patterns(".DS_Store"))
    raw = DEP / "Raw_Data"
    sanger = raw / "01_JH4_Sanger"
    for m in SANGER_MICE:
        shutil.copytree(AB1 / m, sanger / "Original_Chromatograms_AB1" / m,
                        ignore=shutil.ignore_patterns(".DS_Store"))
    shutil.copy(CAL / "samplesheets/jh4_manuscript.csv", sanger / "JH4_Sanger_sample_sheet.csv")
    shutil.copytree(HERE / "sanger_calls", sanger / "SHMtrace_Sanger_calls")
    out = raw / "05_SHMtrace_Analysis"
    for src, name in RUNS.items():
        shutil.copytree(CAL / "figures" / src, out / name, ignore=shutil.ignore_patterns(".DS_Store"))
    shutil.copytree(HERE / "jh4_calls", out / "JH4_529bp_window/calls_by_sample")
    shutil.copytree(HERE / "n114_calls", out / "N114_V_region/calls_by_sample")
    (out / "README.md").write_text(DEPOSIT_ANALYSIS_README)
    (DEP / "README.md").write_text(DEPOSIT_README)
    # manifest of what Ver10 adds, and checksums over every data file
    man = raw / "04_Manifests_and_Integrity"
    added = sorted(p for p in raw.rglob("*") if p.is_file() and
                   (p.is_relative_to(out) or p.is_relative_to(sanger / "Original_Chromatograms_AB1")
                    or p.is_relative_to(sanger / "SHMtrace_Sanger_calls")
                    or p.name == "JH4_Sanger_sample_sheet.csv"))
    with open(man / "VER10_ADDITIONS_MANIFEST.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["path", "bytes", "sha256"])
        for p in added:
            w.writerow([str(p.relative_to(raw)), p.stat().st_size, sha256(p)])
    sums = man / "SHA256SUMS.txt"
    files = sorted(p for p in raw.rglob("*") if p.is_file() and p.parent != man)
    sums.write_text("".join(f"{sha256(p)}  {p.relative_to(raw)}\n" for p in files))
    print(f"deposit: {DEP.name}, {len(files)} checksummed files, {len(added)} added in Ver10")


def build_package():
    if PKG.exists():
        shutil.rmtree(PKG)
    for d in ("1_Manuscript", "2_Cover_Letter", "3_Highlights", "4_Figures/PNG_600dpi", "4_Figures/TIFF_600dpi",
              "4_Figures/Vector_PDF", "5_Graphical_Abstract", "6_Code", "7_Data", "8_Supporting"):
        (PKG / d).mkdir(parents=True)
    ms = PKG / "1_Manuscript/ONT_SHM_Manuscript_JIM_Ver18.docx"
    shutil.copy(MS, ms)
    to_pdf(ms, ms.parent)
    (ms.parent / "README.md").write_text(
        "# Manuscript\n\n`ONT_SHM_Manuscript_JIM_Ver18.docx` and its PDF. Passages changed since Ver16 are "
        "highlighted yellow; remove the highlighting before submission. See `8_Supporting/CHANGES_Ver18.md`.\n")

    cover = Document(V13 / "2_Cover_Letter/ONT_SHM_CoverLetter_JIM.docx")
    for i, text in COVER.items():
        set_paragraph(cover.paragraphs[i], text)
    cl = PKG / "2_Cover_Letter/ONT_SHM_CoverLetter_JIM.docx"
    cover.save(cl)
    to_pdf(cl, cl.parent)
    hl = Document(V13 / "3_Highlights/ONT_SHM_Highlights.docx")
    for i, text in HIGHLIGHTS.items():
        set_paragraph(hl.paragraphs[i], text)
        if len(text) > 85:
            raise RuntimeError(f"highlight over 85 characters: {text}")
    hp = PKG / "3_Highlights/ONT_SHM_Highlights.docx"
    hl.save(hp)
    to_pdf(hp, hp.parent)

    rows = []
    for name, stem in FIGURES.items():
        png = stem.with_suffix(".png")
        shutil.copy(png, PKG / f"4_Figures/PNG_600dpi/{name}.png")
        im = Image.open(png).convert("RGB")
        im.save(PKG / f"4_Figures/TIFF_600dpi/{name}.tif", dpi=(600, 600), compression="tiff_lzw")
        pdf = stem.with_suffix(".pdf")
        if pdf.exists():
            shutil.copy(pdf, PKG / f"4_Figures/Vector_PDF/{name}.pdf")
        rows.append(f"| {name} | `{png.relative_to(REPO)}` | {im.width} × {im.height} |")
    (PKG / "4_Figures/README.md").write_text(
        "# Figures\n\nPNG and LZW-compressed TIFF at 600 dpi, and vector PDF, for every figure in Ver18. Each is "
        "regenerated by the script named in `6_Code/README.md` from the deposited data.\n\n"
        "| Figure | Source in the SHMtrace repository | Pixels |\n|---|---|---|\n" + "\n".join(rows) + "\n")
    shutil.copytree(FIG / "graphical_abstract", PKG / "5_Graphical_Abstract", dirs_exist_ok=True)

    code = PKG / "6_Code"
    shutil.copytree(HERE / "scripts", code / "figure_scripts",
                    ignore=shutil.ignore_patterns("__pycache__", ".DS_Store"))
    (code / "analysis_scripts").mkdir()
    for f in ("jh4_shmtrace.py", "n114_shmtrace.py", "depth_shmtrace.py", "sanger_vs_nanopore.py",
              "sanger_window_figure.py", "organize_jh4_ab1.py", "samplesheets/make_jh4_samplesheets.py"):
        shutil.copy(CAL / f, code / "analysis_scripts")
    for f in ("manuscript_stats.py", "export_sanger_calls.py", "revise_manuscript_sanger.py",
              "revise_manuscript_v18.py", "build_submission_v18.py"):
        shutil.copy(HERE / f, code / "analysis_scripts")
    src = code / "SHMtrace_v0.2_source.zip"
    subprocess.run(["git", "-C", str(REPO), "archive", "--format=zip", "-o", str(src), "HEAD",
                    "src", "pyproject.toml", "README.md", "LICENSE", "tests"], check=True)
    (code / "README.md").write_text(CODE_README)

    data = PKG / "7_Data"
    for src_run, name in RUNS.items():
        d = CAL / "figures" / src_run
        for f in ("samples.csv", "variants.csv", "animal_tests.csv", "table_s1.json", "stats.json",
                  "jh4.json", "n114.json"):
            if (d / f).exists():
                shutil.copy(d / f, data / f"{name}__{f}")
    shutil.copy(FIG / "Individual_mutation_matrices_counts.csv", data)
    (data / "README.md").write_text(
        "# Data underlying the figures\n\nPer-run SHMtrace tables (`<run>__samples.csv`, `__variants.csv`, "
        "`__animal_tests.csv`, `__table_s1.json`) for the JH4 529-bp window, the JH4 full-length interval, the "
        "delivered-read sensitivity run, N114, the depth series and the Sanger–ONT comparison. The complete "
        "project files and reads are in the data deposit (`Raw_Data/05_SHMtrace_Analysis/`).\n")
    shutil.copy(HERE / "CHANGES_Ver18.md", PKG / "8_Supporting/CHANGES_Ver18.md")
    (PKG / "README.md").write_text(PACKAGE_README)
    print(f"package: {PKG.name}")


PACKAGE_README = f"""# Submission package, Ver18 ({STAMP})

**{TITLE}** (Aderinto and Martin). Target: Journal of Immunological Methods.

| Folder | Contents |
|---|---|
| `1_Manuscript/` | Ver18 manuscript (docx, PDF); changes since Ver16 highlighted |
| `2_Cover_Letter/` | Cover letter, updated to the Ver18 results |
| `3_Highlights/` | Five highlights (≤ 85 characters each) |
| `4_Figures/` | Figures 1–4 and S1–S9: PNG and TIFF at 600 dpi, vector PDF |
| `5_Graphical_Abstract/` | Graphical abstract (PNG, TIFF, PDF) |
| `6_Code/` | SHMtrace v0.2 source, the analysis drivers, and the figure scripts |
| `7_Data/` | The tables behind every figure |
| `8_Supporting/` | `CHANGES_Ver18.md`: what changed since Ver16 and why |

Reads, chromatograms and full analysis outputs: `ONT_SHM_DataDeposit_Ver10_{STAMP}`.
"""

DEPOSIT_README = f"""# Raw and intermediate sequencing data, Ver10 ({STAMP})

Companion data deposit for **"{TITLE}"** (Aderinto and Martin, Department of Immunology,
University of Toronto). Intended for SRA/BioProject and Zenodo; the resulting accessions
should be cited in the manuscript's Data and code availability section.

## Contents

| Directory | Contents |
|---|---|
| `Raw_Data/01_JH4_Sanger/` | Sanger clone records (GenBank), the original `.ab1` chromatograms of the six manuscript mice (`Original_Chromatograms_AB1/`), the clone sample sheet, and SHMtrace's Sanger calls |
| `Raw_Data/02_JH4_ONT/` | J<sub>H</sub>4 ONT FASTQs, target-selected and downsampled reads, alignments, the original Geneious exports |
| `Raw_Data/03_N114_ONT/` | N114 V-region equivalents |
| `Raw_Data/04_Manifests_and_Integrity/` | Sample and file manifests, validation, `VER10_ADDITIONS_MANIFEST.csv`, SHA-256 checksums |
| `Raw_Data/05_SHMtrace_Analysis/` | Every SHMtrace analysis behind the Ver18 figures and tables (see its README) |

## What changed from Ver9

- The ONT results in the manuscript are SHMtrace calls (a quality-calibrated test of every candidate against
  the assay's negative controls, false-discovery rate 0.05), made on the deposited 500-read subsets
  (`*_sub500.fastq`). The original Geneious exports are kept.
- The Sanger results come from the `.ab1` chromatograms, now included, rather than the GenBank records.
- `SHA256SUMS.txt` covers every file under `Raw_Data/` except the manifests folder.

## Integrity

```bash
cd Raw_Data && shasum -a 256 -c 04_Manifests_and_Integrity/SHA256SUMS.txt
```
"""

DEPOSIT_ANALYSIS_README = """# SHMtrace analyses (SHMtrace v0.2)

Each folder is one SHMtrace project: `project.json` (every result), `report.html`, CSV tables and
figures. `table_s1.json` gives reads and candidates tested per library (Supplementary Table S1).

| Folder | Input | Used for |
|---|---|---|
| `JH4_529bp_window/` | `02_JH4_ONT/Processed_FASTQ_Depth_Subsets/*_C57BL6_sub500.fastq`, positions 101–629 | Figures 2A, 3; S3; Tables S1, S2 |
| `JH4_full_length_90-1230/` | the same reads, positions 90–1,230 | Figure 2C; S1 |
| `JH4_delivered_reads_sensitivity/` | `02_JH4_ONT/Original_FASTQ/`, SHMtrace's own target selection | sensitivity check (not in the manuscript) |
| `N114_V_region/` | `03_N114_ONT/Processed_FASTQ_Depth_Subsets/*_sub500.fastq`, positions 21–635 | Figure 4; S4, S5 |
| `Depth_series/` | the deposited `_sub30` … `_sub1000` subsets, one project per depth | S6, S7 |
| `Sanger_vs_ONT/` | the chromatograms and the JH4 sub500 reads | Figure 3 statistics, S9 |

Calling: recommended mode (`error_model: qcal`), negative controls Aicda−/− (JH4) and empty vector
(N114); each control is called against the other two. `calibration` in `project.json` records the
leave-one-control-out factor and its trace. Reproduce with `6_Code/analysis_scripts/` of the
submission package.
"""

CODE_README = """# Code

- `SHMtrace_v0.2_source.zip`: the analysis software (Python package `shmtrace`), with its tests.
  Install with `pip install .` from the unpacked folder (Python ≥ 3.11).
- `analysis_scripts/`: the drivers that produced every result from the data deposit.
  - `jh4_shmtrace.py` (and `--full-length`, `--input raw`), `n114_shmtrace.py`, `depth_shmtrace.py`
  - `sanger_vs_nanopore.py --ont-mode recommended --sanger-dir <AB1 folders> --ont-dir <JH4 depth subsets>`
  - `organize_jh4_ab1.py --original`, `make_jh4_samplesheets.py`, `export_sanger_calls.py`: Sanger clones
  - `manuscript_stats.py`: every statistic quoted in the text
  - `revise_manuscript_sanger.py` (Ver16 → Ver17), `revise_manuscript_v18.py` (Ver17 → Ver18),
    `build_submission_v18.py` (this package)
- `figure_scripts/`: one script per figure; `config.py` holds the paths.

Paths at the top of each driver point at the data deposit; edit them to your copy.
"""

if __name__ == "__main__":
    build_deposit()
    build_package()
