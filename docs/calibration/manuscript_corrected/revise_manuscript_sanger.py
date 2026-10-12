"""Update the manuscript for the corrected Sanger data.

    python revise_manuscript_sanger.py SOURCE.docx OUTPUT.docx      (needs python-docx)

SOURCE is never modified. Every replaced or added passage is highlighted yellow so it
can be reviewed in Word. Figures 2, 3 and Supplementary S2 are replaced with the
regenerated ones (aspect ratio kept); Supplementary Figures S8-S9 are added. TOST is removed. Every
edit names the exact text it replaces and the script stops if any is not found.
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
SV = HERE.parent / "figures/sanger_vs_nanopore"

# (paragraph starts with, old text, new text)
EDITS = [
    ("Somatic hypermutation (SHM) diversifies",
     "concordant with matched Sanger data (Pearson r = 0.97)",
     "concordant with matched Sanger data (Pearson r = 0.96)"),
    ("Abbreviations: SHM, somatic hypermutation",
     "; CI, confidence interval; TOST, two one-sided tests.", "; CI, confidence interval."),
    ("Sanger reads were processed with the same",
     "Sanger reads were processed with the same position–substitution definition and Geneious p-value threshold.",
     "Sanger clones were re-analysed from their original .ab1 chromatograms with nanoshm v0.2 in Sanger mode: "
     "each clone was trimmed to its clean, high-quality core and aligned to the C57BL/6 JH4 reference, and every "
     "substitution with base quality ≥ 20 was counted after removing positions at which ≥ 50% of a mouse's clones "
     "differed from the reference. The same position–substitution definition was used."),
    ("Sanger clones were re-analysed from their original",
     "Direct Sanger–ONT comparisons were restricted to the 529-bp JH4 window over which both platforms aligned cleanly and unambiguously.",
     "Direct Sanger–ONT comparisons were restricted to the 529-bp JH4 window over which both platforms aligned cleanly "
     "and unambiguously. This window (positions 101–629) is the JH4 intron, beginning at the JH4 splice donor (the "
     "coding segment ends at position 104), and ends before Sanger read quality declines (Supplementary Figure S8)."),
    ("Depth-standardised unique-variant densities were log10",
     "One Aicda−/− Sanger sample contained no detected call and was assigned the platform's one-event resolution only for log-scale display, summary, and testing. This is a left-censoring convention, not an observed mutation; although it reduces the estimated fold separation, its effects on variance, p-value, and interval width are not guaranteed to be conservative.",
     "After re-analysis every sample contained at least one call, so no left-censoring was required. Sanger "
     "densities divide distinct calls by the bases each clone read within the window, which equals clones × 529 bp "
     "for full-length clones."),
    ("For paired Sanger–ONT spectrum analysis",
     "Platform concordance is therefore described principally from paired profiles and percentage-point differences.",
     "Platform concordance is therefore described principally from paired profiles and percentage-point differences. "
     "Bars in Figure 3B and 3D show each platform's pooled proportion across the three mice with a 95% Wilson "
     "interval, because a mean of per-mouse proportions is dominated by a mouse with few calls."),
    ("For paired Sanger–ONT spectrum analysis",
     " Exploratory paired two one-sided tests (TOST) used a prespecified ±5-percentage-point interpretability margin that was not a validated assay-acceptance criterion. Meeting this criterion for individual types does not establish equivalence of the complete spectrum.",
     ""),
    ("In these samples, Wt mice had markedly higher",
     "The exploratory Welch comparisons were p = 0.0037 for ONT and p = 0.0045 for Sanger. Geometric means were ~830 × 10−5 by Sanger and ~59 × 10−5 by ONT in Wt mice, and ~20 versus ~1.2 × 10−5, respectively, in Aicda−/− controls.",
     "The exploratory Welch comparisons were p = 0.0037 for ONT (~47-fold) and p = 0.0008 for Sanger (~39-fold). "
     "Geometric means were ~910 × 10−5 by Sanger and ~59 × 10−5 by ONT in Wt mice, and ~23 versus ~1.2 × 10−5, "
     "respectively, in Aicda−/− controls."),
    ("The matched ONT and Sanger JH4 datasets",
     "Pooled twelve-type proportions had a descriptive Pearson r = 0.97 and slope = 1.11.",
     "Pooled twelve-type proportions had a descriptive Pearson r = 0.96 and slope = 1.04."),
    ("The matched ONT and Sanger JH4 datasets",
     "(p = 0.875)", "(p = 0.80)"),
    ("The matched ONT and Sanger JH4 datasets",
     "(smallest adjusted p = 0.106)", "(smallest adjusted p = 0.09)"),
    ("The matched ONT and Sanger JH4 datasets",
     " At the exploratory ±5-percentage-point TOST margin, 5 of 12 types met the per-type criterion and 7 were inconclusive (Supplementary Table S2), which does not establish equivalence of the complete spectrum.",
     ""),
    ("The matched ONT and Sanger JH4 datasets",
     "G:C transition and transversion proportions were also descriptively similar (Figure 3D).",
     "G:C transition and transversion proportions were also descriptively similar (Figure 3D). The platforms also "
     "agreed on where calls fell along the intron (25-bp bins; Pearson r = 0.86) and on hotspot preference (G:C "
     "calls at WRC/GYW: Sanger 57%, ONT 55%; A:T calls at WA/TW: 70% and 69%; Supplementary Figure S9)."),
    ("This proof-of-concept study shows that Oxford Nanopore",
     "(Pearson r = 0.97)", "(Pearson r = 0.96)"),
    ("Supplementary Table S1. Per-sample sequencing",
     "JH4 densities use 529 analysed bp;",
     "JH4 densities use 529 analysed bp (for Sanger, the bases each clone read within them); Sanger clones were "
     "re-analysed from the original chromatograms and kept with the mouse they were originally filed under;"),
    ("Supplementary Table S2. Exploratory statistical summary",
     "paired multiplicity-adjusted contrasts, exploratory paired TOST, and a pooled descriptive Pearson correlation",
     "paired multiplicity-adjusted contrasts, and a pooled descriptive Pearson correlation"),
    ("Group geometric means with descriptive 95%",
     "Wt Sanger 828 [492–1192];", "Wt Sanger 908 [764–1183];"),
    ("Group geometric means with descriptive 95%",
     "Aicda−/− Sanger 20 [19–21];", "Aicda−/− Sanger 23 [17–38];"),
    ("Figure 2. Depth-standardised unique-variant density",
     "One Aicda−/− Sanger mouse had zero detected calls and was assigned the one-event scale (18.9 × 10−5) only for log display, summary, and testing. This left-censoring convention reduces the fold difference, but its effects on variance, p-value, and interval width are not guaranteed to be conservative; the Sanger inferential result is therefore descriptive.",
     "Sanger densities divide distinct calls by the bases each clone read within the window, from clones "
     "re-analysed from the original chromatograms; every mouse had at least one call."),
    ("Figure 2. Depth-standardised unique-variant density",
     "~41-fold (Sanger; p = 0.0045)", "~39-fold (Sanger; p = 0.0008)"),
    ("Figure 2. Depth-standardised unique-variant density",
     "Wt Sanger 828 [492–1192], and Aicda−/− Sanger 20 [19–21]",
     "Wt Sanger 908 [764–1183], and Aicda−/− Sanger 23 [17–38]"),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "(n = 515 Sanger and 571 ONT", "(n = 489 Sanger and 571 ONT"),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "bars show mean ± SD and points show mice.",
     "bars show each platform's pooled proportion with a 95% Wilson interval, and points show mice (area "
     "proportional to the mouse's number of calls)."),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "interaction was p = 0.875", "interaction was p = 0.80"),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "(Pearson r = 0.97; slope = 1.11;", "(Pearson r = 0.96; slope = 1.04;"),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "(D) G:C transition and transversion proportions; adjusted paired p = 0.803 and 0.982.",
     "(D) G:C transition and transversion proportions (pooled, 95% Wilson intervals); adjusted paired p = 0.62 and "
     "0.99."),
    ("Figure 3. ONT yields a JH4 substitution spectrum",
     "(E) Calls per sequence (Sanger, n = 111; ONT, n = 1,500 reads). Ninety-two Sanger sequences aligned unambiguously to C57BL/6 for spectrum analysis; 19 carrying residual background haplotypes were retained in the sequence denominator and zero-call class but contributed no calls to the pooled spectrum.",
     "(E) Calls per sequence (Sanger, n = 103 clones; ONT, n = 1,500 reads). Agouti-type sequences, which carry "
     "no SHM once their own background is removed, are excluded from the Sanger set."),
    ("Supplementary Figure S2. Individual JH4 intron",
     "(n = 189, 227, and 99 unique variants, respectively)", "(n = 153, 230, and 106 unique variants, respectively)"),
    ("Supplementary Figure S2. Individual JH4 intron",
     "(B) Aicda−/− mice 1–3 (n = 0, 1, and 1).", "(B) Aicda−/− mice 1–3 (n = 1, 1, and 2)."),
    ("Supplementary Figure S2. Individual JH4 intron",
     "Variants were restricted to single-nucleotide substitutions with variant frequency < 50% and Geneious approximate variant p ≤ 0.05.",
     "Calls are single-nucleotide substitutions with base quality ≥ 20 in the quality-trimmed clone, after removing "
     "positions at which ≥ 50% of a mouse's clones differ from the reference (nanoshm v0.2, Sanger mode), from "
     "clones re-analysed from the original chromatograms."),
    ("Supplementary Figure S2. Individual JH4 intron",
     " Percentages are not defined for the sample with no passing variants.", ""),
]

# (table index, row, column, old, new)
CELLS = [
    (0, 4, 1, "36 to 38 (Wt), 9 to 10 (Aicda−/−)", "26 to 39 (Wt), 10 to 11 (Aicda−/−)"),
    (2, 2, 2, "p = 0.0045 (~41-fold)", "p = 0.0008 (~39-fold)"),
    (2, 4, 2, "r = 0.97 [0.91–0.99]; slope = 1.11", "r = 0.96 [0.86–0.99]; slope = 1.04"),
    (2, 5, 2, "interaction p = 0.875; none significant (minimum adjusted p = 0.106)",
     "interaction p = 0.80; none significant (minimum adjusted p = 0.09)"),
    (2, 7, 2, "Ts adjusted p = 0.803; Tv adjusted p = 0.982", "Ts adjusted p = 0.62; Tv adjusted p = 0.99"),
]
# table rows removed (table index, row, text the row must contain)
DROP_ROWS = [(2, 6, "TOST")]
# Supplementary Table S1 Sanger rows: clones, bp, candidates, retained, density (x1e-5)
S1_ROWS = {"W2": ("39", "529", "153", "153", "764"), "W3": ("38", "529", "230", "230", "1183"),
           "W4": ("26", "529", "106", "106", "828"), "AID1": ("11", "529", "1", "1", "17.3"),
           "AID2": ("10", "529", "1", "1", "19.0"), "AID3": ("10", "529", "2", "2", "37.8")}

NEW_METHODS = (
    "Sanger clone assignment. The 447 chromatograms (eight sequencing orders) were matched to the hand-collated "
    "GenBank clone records, and each clone was kept with the mouse under which it was originally filed; no clone "
    "was reassigned between mice. Clones labelled \"jh6\", from a second primer set, aligned across the JH4 intron and were "
    "included. Two GenBank records filed with Wt2 were renamed copies of MKLN-knockout (K4) chromatograms "
    "and were excluded, as was one clone filed with Wt4 that was identical over 520 bp, including all seven "
    "of its mutations, to two MKLN-knockout clones sequenced on the same plate and in a separate order. "
    "Sequences matching the Agouti-background JH4-like reference (~12% divergent from C57BL/6) were "
    "excluded; aligned to that reference with its own background removed, they carried no SHM. Chromatograms "
    "without a usable high-quality segment or alignment were excluded, and two clones never filed with any mouse "
    "were not used.")

NEW_FIGS = [
    ("Supplementary Figure S8. ", SV / "figS_sanger_window_original.png",
     "The 529-bp comparison window is the JH4 intron and the region the Sanger clones read cleanly. Top: share of "
     "Wt Sanger clones with a base of quality ≥ 20 at each JH4 reference position. Bottom: distinct Wt calls per "
     "10³ bases read, 25-bp bins. Blue shading, the 529-bp window (positions 101–629); orange, the 27-bp extensions "
     "of two 556-bp alternatives (74–100, the end of the JH4 coding segment; 630–656, where the forward reads fade); "
     "dashed line, end of the JH4 coding segment (position 104). Sanger background in Aicda−/− clones was 5 "
     "mismatches in 16,344 bases inside the window. Either alternative changed the pooled Wt call count by about 3%."),
    ("Supplementary Figure S9. ", SV / "geneious_original/fig3_agreement_position_hotspot.png",
     "Sanger and ONT agree beyond the substitution spectrum. ONT calls use the manuscript's Geneious rule "
     "(reproduced in nanoshm); calls from the three Wt mice pooled on each platform. (A) Share of distinct calls "
     "along the 529-bp window in 25-bp bins (Pearson r = 0.86). (B) Share of G:C calls at WRC/GYW and of A:T calls "
     "at WA/TW (bars, mean of mice; points, mice), against the share of such sites in the window (dashed)."),
]


def replace_text(paragraph, old, new):
    runs = paragraph.runs
    full = "".join(r.text for r in runs)
    if full != paragraph.text:
        raise RuntimeError(f"paragraph has non-run content: {paragraph.text[:60]!r}")
    i = full.find(old)
    if i < 0:
        raise RuntimeError(f"not found: {old[:70]!r} in {full[:60]!r}")
    j = i + len(old)
    segs, pos = [], 0
    for r in runs:
        segs.append((pos, pos + len(r.text), r))
        pos += len(r.text)
    first = next(r for a, b, r in segs if a <= i < b) if old else runs[0]
    new_runs = []
    for a, b, r in segs:
        if b <= i or a >= j:
            new_runs.append((r.text, r))
            continue
        if a < i:
            new_runs.append((r.text[: i - a], r))
        if r is first:
            new_runs.append((new, None))
        if b > j:
            new_runs.append((r.text[j - a:], r))
    p = paragraph._p
    for r in runs:
        p.remove(r._r)
    for text, src in new_runs:
        if not text:
            continue
        base = src if src is not None else first
        el = copy.deepcopy(base._r)
        p.append(el)
        from docx.text.run import Run
        run = Run(el, paragraph)
        run.text = text
        if src is None:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def set_cell(cell, old, new):
    for par in cell.paragraphs:
        if old in par.text:
            replace_text(par, old, new)
            return
    raise RuntimeError(f"cell text not found: {old!r} in {cell.text!r}")


def swap_image(doc, paragraph, path):
    blip = next(paragraph._p.iter(qn("a:blip")))
    part = doc.part.related_parts[blip.get(qn("r:embed"))]
    part._blob = path.read_bytes()
    w, h = Image.open(path).size
    ext = paragraph._p.find(".//" + qn("wp:extent"))
    cx = int(ext.get("cx"))
    cy = int(cx * h / w)
    ext.set("cy", str(cy))
    for a_ext in paragraph._p.iter(qn("a:ext")):
        if a_ext.get("cx") == str(cx):
            a_ext.set("cy", str(cy))


def main(src, out):
    doc = Document(src)
    paras = doc.paragraphs

    def find(start):
        hits = [p for p in paras if p.text.startswith(start)]
        if len(hits) != 1:
            raise RuntimeError(f"{len(hits)} paragraphs start with {start!r}")
        return hits[0]

    for start, old, new in EDITS:
        replace_text(find(start), old, new)
    for t, r, c, old, new in CELLS:
        set_cell(doc.tables[t].rows[r].cells[c], old, new)
    for t, r, must in DROP_ROWS:
        row = doc.tables[t].rows[r]
        if must not in "".join(c.text for c in row.cells):
            raise RuntimeError(f"table {t} row {r} lacks {must!r}")
        row._tr.getparent().remove(row._tr)
    s1 = doc.tables[1]
    for row in s1.rows:
        key = row.cells[0].text.strip()
        if key in S1_ROWS and row.cells[2].text.strip() == "Sanger":
            for col, val in zip((5, 6, 7, 8, 9), S1_ROWS[key], strict=True):
                cell = row.cells[col]
                if cell.text.strip() != val:
                    set_cell(cell, cell.text.strip(), val)

    # a Methods paragraph on clone curation, after the Sanger processing paragraph
    anchor = find("Sanger clones were re-analysed from their original")
    new_p = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    np_ = Paragraph(new_p, anchor._parent)
    for r in list(np_.runs)[1:]:
        new_p.remove(r._r)
    np_.runs[0].text = NEW_METHODS
    np_.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW

    # regenerated figures
    for start, path in (("Figure 2. Depth-standardised", FIG / "Figure2_JH4_combined.png"),
                        ("Figure 3. ONT yields", FIG / "Figure3_concordance.png"),
                        ("Supplementary Figure S2. Individual", FIG / "FigureS5_individual_Sanger_matrices.png")):
        legend = find(start)
        idx = [p._p for p in doc.paragraphs].index(legend._p)
        img_p = next(doc.paragraphs[k] for k in range(idx - 1, idx - 4, -1)
                     if list(doc.paragraphs[k]._p.iter(qn("a:blip"))))
        swap_image(doc, img_p, path)

    # new supplementary figures, after S7, in the same legend style
    model = find("Supplementary Figure S7.")
    tail = model._p
    for label, path, text in NEW_FIGS:
        pic = copy.deepcopy(model._p)
        for r in list(pic):
            if r.tag != qn("w:pPr"):
                pic.remove(r)
        tail.addnext(pic)
        pp = Paragraph(pic, model._parent)
        pp.add_run().add_picture(str(path), width=Inches(6.4))
        leg = copy.deepcopy(model._p)
        pic.addnext(leg)
        lp = Paragraph(leg, model._parent)
        runs = lp.runs
        bold_run, body_run = runs[0], runs[1] if len(runs) > 1 else runs[0]
        for r in runs:
            if r is not bold_run and r is not body_run:
                leg.remove(r._r)
        bold_run.text = label
        body_run.text = text
        for r in (bold_run, body_run):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        tail = leg
    doc.core_properties.comments = ("Ver17: Sanger results re-analysed from the chromatograms, clones as originally "
                                    "filed (nanoshm docs/calibration/manuscript_corrected)")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
