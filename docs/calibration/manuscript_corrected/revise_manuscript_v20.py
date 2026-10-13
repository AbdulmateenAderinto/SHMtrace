"""Ver20: Figure 1 redrawn in SciCanvas as an overview; the detailed workflow as Supplementary Figure S11.

    python revise_manuscript_v20.py VER19.docx VER20.docx      (needs python-docx)

Figures: figure1_scicanvas/figure1_overview.js and figure1_detailed.js, built with
figure1_scicanvas/build.js (SciCanvas's own renderer). The two in-text citations of Figure 1 also
cite S11, and the AI declaration says how the artwork was made. New text is highlighted yellow.
SOURCE is never modified, and the script stops if an anchor is not found.
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

from revise_manuscript_sanger import swap_image
from revise_manuscript_v18 import replace_text, restyle

HERE = Path(__file__).resolve().parent
OUT = HERE / "figure1_scicanvas/out"
OVERVIEW = OUT / "Figure1_overview_preview.png"           # 300 dpi in the docx; production files: PDF, 1,000-dpi TIFF
DETAILED = OUT / "Figure_workflow_detailed_preview.png"

FIG1_LEGEND = (
    "Study overview. (1) Samples. The JH4 downstream intron (1.23 kb) was amplified from Peyer's patch "
    "germinal-centre B cells of Wt and Aicda−/− mice (n = 3 per genotype), and a fixed V region (0.65 kb) from N114 "
    "hybridoma subclones transduced with pMX-Aicda or the empty vector (three subclones per vector). Aicda−/− mice "
    "and empty-vector subclones are the AID-null controls. Amplicons are drawn to scale. (2) Sequencing. Full-length "
    "amplicons were sequenced by ONT (R10.4.1, Kit 14, SUP basecalling) without cloning and analysed at 500 reads per "
    "library; JH4 amplicons from the same mice were cloned and Sanger sequenced in separate PCRs as a benchmark. "
    "(3) Analysis with SHMtrace. Reads from a related gene copy, joined (chimeric) reads and low-quality reads were "
    "removed; a change was kept only if it appeared more often than sequencing errors do in the AID-null controls, "
    "with at most 5% of kept calls expected to be errors; and groups were compared with one value per mouse or "
    "subclone. Sanger clones took the same steps, keeping only confident base calls. Cases are shown in blue and "
    "AID-null controls in red, as in Figures 2 and 4. (4) Readouts, with the "
    "figures that show them. The full workflow, with read numbers at each step, is in Supplementary Figure S11. "
    "Drawn in SciCanvas; no generative image model was used.")

S11_LEGEND = (
    "Detailed workflow. (A) Experimental design. JH4: Peyer's patch B220+GL7+Fas+ germinal-centre B cells were "
    "sorted from 4-month-old Wt and Aicda−/− mice (n = 3 per genotype), and the 1,230-bp JH4 intron amplicon was "
    "amplified with SuperFi II for ONT and, in separate PCRs, cloned for Sanger sequencing. ONT analysis covers "
    "positions 90–1,230 (1,141 bp); the window shared with Sanger is positions 101–629 (529 bp). N114: cells were "
    "transduced with pMX-Aicda or the empty vector (both GFP and puromycin resistance), sorted for GFP-high "
    "expression, subcloned by limiting dilution and kept under puromycin for 1 month; three subclones per vector were "
    "sequenced, and positions 21–635 of the 655-bp reference (615 bp) were analysed. Amplicons are drawn to scale. "
    "(B) Read preprocessing: (1) reads matching the target better than a competing reference were kept (JH4: "
    "C57BL/6 against Agouti, with a recurrent-contaminant screen; N114: one reference); (2) full-length reads were "
    "kept (JH4 900–2,000 bp, N114 400–900 bp); (3) reads averaging more than 0.01 expected errors per base were "
    "removed; (4) each library was down-sampled once to 500 reads, and the subsets are deposited. (C) Finding "
    "mutations with SHMtrace: (5) reads were aligned to the reference, and reads matching a related gene copy (ten "
    "or more of its distinguishing bases) were dropped; (6) every base that differs from the reference was listed, a "
    "change in one read being enough; (7) how often sequencing errors produce each type of change was measured in "
    "the AID-null controls, by strand and base-call quality, with a five-fold safety margin set by testing each "
    "control against the others; (8) changes seen in more reads than errors explain were kept, with at most 5% of "
    "kept calls expected to be errors (Benjamini–Hochberg false-discovery rate), and changes present in half or "
    "more of the reads were dropped as inherited rather than somatic; (9) results are summarised per animal "
    "(distinct mutations per base read, spectra, clones and trees), each mouse or subclone counting once. (D) Sanger "
    "clones take the same route: 447 chromatograms from eight sequencing orders were assigned to mice with a sample "
    "sheet and quality-control checks, trimmed of poor-quality ends, aligned to JH4, and called only where the base "
    "call is at least 99% accurate (Q ≥ 20). Cases are blue and AID-null controls red. (E) Reads delivered per "
    "library (grey) and reads passing steps 1–3 (coloured); dashed line, the 500 reads analysed. Drawn in SciCanvas; "
    "no generative image model was used.")

CITES = [("Raw reads first underwent assay-specific target selection", "(Figure 1)", "(Figure 1; Supplementary Figure S11)"),
         ("To test whether ONT could recover", "(Figure 1)", "(Figure 1; Supplementary Figure S11)")]
# "Figure 1" in this paragraph is a cross-reference link; only the plain text after it is replaced
AI_OLD = (" is explanatory vector artwork generated programmatically in Matplotlib from author-specified "
          "content; Codex assisted with layout planning and code drafting, but no generative image model produced or "
          "altered the artwork.")
AI_NEW = (" and Supplementary Figure S11 are explanatory artwork drawn in SciCanvas from author-specified "
          "content, laid out by code drafted with Anthropic Claude; no generative image model produced or altered the "
          "artwork.")


def main(src, out):
    doc = Document(src)

    def find(start):
        hits = [p for p in doc.paragraphs if p.text.startswith(start)]
        if len(hits) != 1:
            raise RuntimeError(f"{len(hits)} paragraphs start with {start!r}")
        return hits[0]

    for start, old, new in CITES:
        replace_text(find(start), old, new)
    replace_text(find("Declaration of Generative AI"), AI_OLD, AI_NEW)

    # Figure 1: new picture (aspect kept by swap_image), new legend body
    leg1 = find("Figure 1. ")
    idx = [p._p for p in doc.paragraphs].index(leg1._p)
    img = next(doc.paragraphs[k] for k in range(idx - 1, idx - 4, -1) if list(doc.paragraphs[k]._p.iter(qn("a:blip"))))
    swap_image(doc, img, OVERVIEW)
    body = "".join(r.text for r in leg1.runs)
    head = body[: body.index("workflow for full-length ONT amplicons.") + len("workflow for full-length ONT amplicons.")]
    replace_text(leg1, head, "Figure 1.")          # the old bold title
    replace_text(leg1, "".join(r.text for r in leg1.runs)[len("Figure 1."):], " " + FIG1_LEGEND)
    for r in leg1.runs:
        if r.text.startswith(" Study overview."):
            r.bold = None
    lead = leg1.runs[0]
    lead.bold = True

    # Supplementary Figure S11 after S10, in S10's style
    s10 = find("Supplementary Figure S10.")
    pic = copy.deepcopy(s10._p)
    for child in list(pic):
        if child.tag.split("}")[1] != "pPr":
            pic.remove(child)
    s10._p.addnext(pic)
    Paragraph(pic, s10._parent).add_run().add_picture(str(DETAILED), width=Inches(6.4))
    leg = copy.deepcopy(s10._p)
    pic.addnext(leg)
    lp = Paragraph(leg, s10._parent)
    bold, body_run = lp.runs[0], lp.runs[1]
    for r in lp.runs[2:]:
        leg.remove(r._r)
    bold.text, body_run.text = "Supplementary Figure S11. ", S11_LEGEND
    for r in (bold, body_run):
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for p in doc.paragraphs:
        restyle(p)
    doc.core_properties.comments = "Ver20: Figure 1 overview and Supplementary Figure S11 (detailed workflow), SciCanvas"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
