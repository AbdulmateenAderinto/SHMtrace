"""Ver21: every ONT result from SHMtrace run end to end on the reads as delivered, from Ver20.

    python revise_manuscript_v21.py VER20.docx VER21.docx      (needs python-docx)

Up to Ver20 the ONT inputs were 500-read subsets made by an earlier preprocessing script
(900-2,000 bp length window, 15-mer competitive selection, a 651-nt N114 selection reference,
reservoir sampling with seed 42). Ver21 reports SHMtrace doing all of it from the delivered
FASTQs: length and expected-error gates, competitive alignment against paralogue references,
the paralogue-allele filter, a random draw of 500 usable reads (seed 0) and the calling.
  JH4  docs/calibration/jh4_shmtrace.py (529 bp) and --full-length (positions 90-1,230)
  N114 docs/calibration/n114_shmtrace.py
  depth series depth_shmtrace.py; Sanger-ONT agreement sanger_vs_nanopore.py (--ont-dir Original_FASTQ)
  read accounting read_accounting.py (table_s1.json per run); trees mutation_tree_figure.py
Methods 2.7 is rewritten to SHMtrace's steps, and every number, table, legend and figure that
changed is replaced. Replaced text is highlighted yellow. SOURCE is never modified, and the
script stops if any expected text is not found.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from revise_manuscript_sanger import set_cell, swap_image
from revise_manuscript_v18 import replace_text, restyle

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
CAL = HERE.parent / "figures"
SV = CAL / "sanger_vs_nanopore/recommended_manuscript"
S1F = HERE / "figure1_scicanvas/out"

# ---------------------------------------------------------------- whole paragraphs (start -> new text)
PARAS = {
    # ---- Abstract
    "Somatic hypermutation (SHM) diversifies": (
        "Somatic hypermutation (SHM) diversifies immunoglobulin genes through activation-induced cytidine deaminase "
        "(AID) but is commonly measured by labour-intensive cloned-amplicon Sanger sequencing. We evaluated a "
        "cloning-free workflow using Oxford Nanopore Technologies (ONT) R10.4.1 flow cells and Kit 14 chemistry to "
        "sequence complete ~1.23-kb JH4 and ~0.65-kb hybridoma V-region amplicons. SHMtrace, written for this study, "
        "analysed the delivered reads end to end: read-length and expected-error gates, competitive alignment against "
        "paralogue references, a random draw of 500 usable reads per library, and a quality-calibrated test of each "
        "candidate substitution against the assay's negative controls at a false-discovery rate of 0.05, yielding an "
        "operational set of filtered position–substitution calls. In three wild-type and three Aicda−/− mice, ONT "
        "recovered the expected group-level AID-dependent difference and a pooled substitution spectrum descriptively "
        "concordant with matched Sanger data (Pearson r = 0.93). Full-length reads extended analysis from the 529-bp "
        "shared window to a reliably covered 1,141-bp interval while retaining a similar pooled spectrum. In three "
        "AID-transduced and three empty-vector N114 subclones, ONT detected at least a 92-fold group difference, with "
        "no calls in any empty-vector control, and the expected G:C-biased signature. At 500 reads per sample, a "
        "single random draw per depth recovered a mean of 91–94% of the filtered positions observed at the deepest "
        "analysed depth; AID-deficient controls retained a post-filter apparent-variant background of "
        "~0.5 × 10−5 per analysed base. Because the workflow lacks unique molecular identifiers and did not "
        "orthogonally confirm each retained position, it supports proof-of-concept, matched-depth group comparisons "
        "of apparent-variant yield and aggregate spectra rather than validation of every rare variant or "
        "molecule-level mutation frequency. Broader use requires locus-, run-, chemistry-, and basecaller-specific "
        "validation."),
    # ---- Methods 2.7, now SHMtrace's read preparation
    "Raw reads first underwent assay-specific target selection": (
        "Reads were prepared by SHMtrace v0.2 (Section 2.8; Data and code availability) directly from the FASTQ files "
        "as delivered by the sequencing provider, with no separate preprocessing (Figure 1; Supplementary Figure S11). "
        "Each read was first gated on raw length, with an upper limit of 2,000 bp for JH4 and 900 bp for N114 to "
        "exclude concatemers and over-length off-target products. No minimum length was applied at this step; "
        "truncated reads are removed by the aligned-length criterion below."),
    "Target-selected JH4 reads were retained": (
        "Reads were next filtered by an expected-error rate calculated from Phred+33 quality scores as the sum of "
        "10^(−Q/10) across the read divided by read length (Edgar and Flyvbjerg, 2015). Reads averaging ≤ 0.01 "
        "expected errors per base were retained. This is a Q20-equivalent average predicted-error threshold, not a "
        "measurement of realised per-base accuracy."),
    "Reads were next filtered by an expected-error rate": (
        "The ≤ 0.01 expected-errors-per-base threshold was used as an operational read-level quality-control cutoff "
        "to reduce, rather than eliminate, the contribution of low-quality sequence. Normalisation by read length "
        "allowed the same average-error criterion to be applied to the two amplicons. The cutoff was selected "
        "empirically for these targets and this basecalling model but was not calibrated against a per-base truth set "
        "and should not be interpreted as a validated sensitivity–specificity optimum. Different loci, sequencing "
        "runs, chemistries, or basecalling models require re-derivation."),
    "The ≤0.01 expected-errors-per-base threshold was used": (
        "Reads passing both gates were aligned with minimap2 (map-ont) to the target reference together with "
        "competing references, and each read was assigned to the sequence it aligned to best. For JH4, the competing "
        "references were a JH4-like paralogue (the Agouti-background sequence, ~12% divergent from C57BL/6) and a "
        "recurrent contaminant haplotype; N114 had a single reference. Reads aligning best to a competing reference "
        "were removed, as were unmapped reads, chimeric reads (those with a supplementary alignment) and reads whose "
        "aligned length was below 800 bp (JH4) or 400 bp (N114). Because a paralogue read longer than the paralogue "
        "record, or a PCR chimera with it, can still align best to the target, JH4 reads carrying ten or more bases "
        "specific to a competing reference were then removed; real JH4 reads carry 0–7 such bases and paralogue reads "
        "and chimeras 11–62. The same references were used for selection and for variant analysis."),
    "For the primary analysis, target-selected reads passing": (
        "Each barcoded library yielded 2,336–5,000 delivered reads, of which 962–2,905 were usable after these steps "
        "(mean 1,620; median 1,456). Across the six JH4 libraries, 9% of delivered reads were removed as over-length, "
        "29% by the expected-error gate and 28% at alignment or by the paralogue-allele filter, leaving 33% usable; "
        "for N114 the corresponding figures were 25%, 22%, 9% and 43% (Supplementary Figure S11E; Supplementary "
        "Table S1). SHMtrace then drew 500 usable reads per library at random with a fixed seed, so 500 reads was a "
        "matched working depth rather than a limit imposed by yield, and the analysed reads are reproducible from the "
        "delivered files. For the depth series, one random draw was made per library at each of 30, 60, 125, 250, "
        "500 and 1,000 reads (or all usable reads where fewer were available); repeated-draw sampling uncertainty was "
        "not assessed. Downsampling does not identify PCR duplicates, which remain sampled in proportion to their "
        "abundance in this non-UMI workflow."),
    # ---- Results
    "In these samples, Wt mice had markedly higher": (
        "In these samples, Wt mice had markedly higher depth-standardised JH4 unique-variant densities than Aicda−/− "
        "mice on both platforms (Figure 2A,B), consistent with the expected AID dependence. The exploratory Welch "
        "comparisons were p = 0.0025 for ONT (≥ 137-fold, the one Aicda−/− library without calls placed at its "
        "one-event density) and p = 0.0009 for Sanger (~38-fold). Geometric means were ~870 × 10−5 by Sanger and "
        "~66 × 10−5 by ONT in Wt mice, and ~23 versus ~0.48 × 10−5, respectively, in Aicda−/− controls. These "
        "operational metrics have different sampling denominators and should be compared within rather than between "
        "platforms; they are not interchangeable per-molecule mutation frequencies."),
    "After all filters, the three Aicda−/− libraries retained": (
        "After all filters, the three Aicda−/− libraries retained 0, 2 and 1 calls, a geometric-mean apparent "
        "background of ~0.48 × 10−5 distinct calls per analysed base (pooled, 3 calls in ~790,000 analysed "
        "read-bases, about one per 260,000). Expressed as Q = −10 × log10(rate), this is a Q53-equivalent "
        "assay-level indicator, not raw ONT accuracy or an individual-call false-positive probability. Residual calls "
        "may include technical error, PCR error, low-frequency haplotypes, contamination, or genuine AID-independent "
        "variation. In the three Wt libraries, 415, 534 and 306 candidate position–substitutions were tested and 205, "
        "340 and 73 retained; in the Aicda−/− libraries, 214, 256 and 254 were tested (Supplementary Table S1). An "
        "earlier analysis of separately preprocessed 500-read subsets of the same libraries in Geneious Prime "
        "2026.1.2 (Kearse et al., 2012), with an unadjusted per-candidate p ≤ 0.05, retained a similar number of Wt "
        "calls (161, 342 and 68) but more Aicda−/− calls (2, 6 and 3; ~1.2 × 10−5) and a smaller group difference "
        "(~47-fold; p = 0.0037). Calling against the controls with false-discovery control therefore mainly reduced "
        "the background."),
    "Because ONT reads span the complete amplicon": (
        "Because ONT reads span the complete amplicon, we profiled filtered apparent variants across the JH4 product "
        "rather than only the 529-bp Sanger-comparison window (Figure 2C). Coverage was uniform (~500×) across "
        "positions 90–1,230 but low and uneven over the first ~90 bp; analysis therefore used the reliably covered "
        "1,141-bp region. Within this interval, Wt samples had a higher depth-standardised density than Aicda−/− "
        "controls (39 versus 0.51 × 10−5; ~78-fold). Approximately 70% of Wt calls fell within the first ~450 bp, "
        "and the number of calls fell about eight-fold between the first and last 250 bp of the interval. Full-length "
        "reads thus revealed a sparser distal apparent-variant landscape beyond the classical window. Pooled "
        "substitution proportions were similar between the 529-bp and 1,141-bp analyses (Pearson r = 0.986; all "
        "differences ≤ 1.8 percentage points), while the longer interval contained more distinct calls (795 versus "
        "618; Supplementary Figure S1). This comparison is descriptive and does not establish candidate-level truth."),
    "Because each read spans the amplicon, the reads of one B-cell clone": (
        "Because each read spans the amplicon, the reads of one B-cell clone can also be ordered into a mutation tree "
        "(Supplementary Figure S10). The largest Wt3 clone (69 reads sharing a V(D)J junction) carried four founder "
        "mutations in every mutated read and then divided into five sub-lineages, up to 17 mutations deep; 70% of its "
        "mutated reads fitted a single tree, and 0.6% of its mutation pairs (7 of 1,176) failed the four-gamete test. "
        "Every other Wt clone with at least 10 reads had no incompatible pairs. In the other Wt mice, the dominant Wt2 "
        "clone (323 reads) was largely unmutated, and each Wt4 clone carried one block of shared mutations without "
        "sub-branches, also when all 962 usable Wt4 reads were analysed. Without UMIs, recombination between "
        "molecules during PCR can create incompatible pairs, so these trees are read-level reconstructions rather "
        "than confirmed genealogies."),
    "The matched ONT and Sanger JH4 datasets produced": (
        "The matched ONT and Sanger JH4 datasets produced similar aggregate substitution profiles (Figure 3A–C). "
        "Pooled twelve-type proportions had a descriptive Pearson r = 0.93 and slope = 1.00. The exploratory "
        "repeated-measures analysis detected no platform-by-type interaction (p = 0.61), and no individual type met "
        "the Šídák-adjusted paired significance criterion (smallest adjusted p = 0.68); these non-significant results "
        "do not establish equality. G:C transition and transversion proportions were also descriptively similar "
        "(Figure 3D). The platforms also agreed on where calls fell along the intron (25-bp bins; Pearson r = 0.92) "
        "and on hotspot preference (G:C calls at WRC/GYW: Sanger 58%, ONT 54%; A:T calls at WA/TW: Sanger 70%, ONT "
        "69%; Supplementary Figure S9). Per-sequence mutation-load distributions differed in shape (Figure 3E), "
        "consistent with different sampling and filtering of bulk ONT reads and cloned Sanger amplicons. Concordance "
        "therefore applies to aggregate category profiles, not one-to-one confirmation of ONT-only rare positions."),
    "At matched 500-read depth, pMX-Aicda subclones": (
        "At matched 500-read depth, pMX-Aicda subclones had a geometric-mean depth-standardised density of "
        "30 × 10−5, whereas none of the three empty-vector controls had a call (Figure 4A). With the controls placed "
        "at their one-event density (0.33 × 10−5), the group difference was at least 92-fold (exploratory Welch "
        "p = 0.0048). The pooled spectrum was dominated by C→T and G→A substitutions at G:C base pairs (96% of calls "
        "at G:C), with few A:T substitutions (Figure 4), consistent with AID-catalysed cytidine deamination (Di Noia "
        "and Neuberger, 2002; Rada et al., 2004; Feng et al., 2021). All three pMX-Aicda subclones showed the "
        "G:C-biased profile, whereas the controls had no calls (Supplementary Figure S4). Read-level recurrence "
        "weighting retained the same pooled pattern (Supplementary Figure S5), but without UMIs these observations "
        "are not counts of independent original molecules."),
    "We evaluated one independently drawn, non-nested reservoir subsample": (
        "We analysed each library at 30, 60, 125, 250, 500 and 1,000 reads, SHMtrace drawing one random subset of its "
        "usable reads at each depth, and re-called variants with the controls calibrated at the same depth "
        "(Supplementary Figures S6 and S7). In AID-active samples, the number of filtered position–substitution calls "
        "rose with depth; at 500 reads, mean counts were 94% (JH4) and 91% (N114) of each sample's maximum. Controls "
        "retained few calls at every depth: 0–3 per JH4 Aicda−/− library, and 0–5 per N114 empty-vector subclone at "
        "30 and 60 reads and none from 125 reads upward. These percentages describe this finite dataset and single "
        "subsampling realisation; they do not estimate recovery of the unknown true mutation repertoire or establish "
        "a universal saturation depth."),
    "Below ~250 reads, a mutation seen in a single read": (
        "Depth-standardised density therefore declined with depth, as expected when its numerator saturates: the Wt "
        "geometric mean fell from ~460 × 10−5 at 30 reads to 66 × 10−5 at 500 and 36 × 10−5 at 1,000 reads, and the "
        "pMX-Aicda mean from ~110 to 30 and 16 × 10−5 (Supplementary Figure S6). Saturation depended on clonal "
        "structure: the clonally shared mutations of Wt3 were nearly all called by 250 reads (332 of a maximum of "
        "340), whereas Wt2, whose mutations are mostly carried by single reads, was still gaining calls between 500 "
        "and 1,000 reads (205 and 247). Greater depth does not improve an individual read's accuracy; it samples more "
        "amplified products and gives each candidate more opportunity to reach significance. Density is therefore "
        "comparable only at matched depth, and cross-platform support comes from the shared group-level AID "
        "dependence and descriptively similar aggregate spectrum, not equality of absolute density or confirmation "
        "of each position."),
    # ---- Conclusions
    "This proof-of-concept study shows that Oxford Nanopore": (
        "This proof-of-concept study shows that Oxford Nanopore R10.4.1/Kit 14 sequencing can support cloning-free, "
        "full-length, substitution-focused SHM profiling at the group and aggregate-spectrum level under the tested "
        "conditions. ONT recovered a JH4 spectrum descriptively concordant with matched Sanger data (Pearson "
        "r = 0.93), a marked AID-dependent group difference in the JH4 intron, and at least a 92-fold difference "
        "between AID-expressing and empty-vector N114 subclones, with no calls in any of the three empty-vector "
        "controls. The post-filter Aicda−/− background was ~0.5 × 10−5 apparent variants per analysed base, and "
        "full-length reads supported analysis of a reliably covered 1,141-bp JH4 interval. These results establish "
        "feasibility but do not validate every low-frequency position, establish independent-molecule mutation "
        "frequencies, or demonstrate transferability across runs and loci. Further prospective validation using "
        "orthogonal truth sets, explicit haplotype controls, independently balanced runs, and UMI or consensus "
        "approaches would be required before broader use for individual rare-variant inference."),
    "Data and code availability:": (
        "Data and code availability: The ONT FASTQ files as delivered by the sequencing provider, which are "
        "SHMtrace's inputs, and the Sanger chromatograms and clone records; reference sequences; SHMtrace project "
        "files, per-library read accounting and variant calls; the 500-read subsets made by an earlier preprocessing "
        "script and the Geneious variant-call exports made from them, retained for comparison; sample and file "
        "manifests; and integrity checksums are provided with this submission for peer review. SHMtrace draws the "
        "analysed reads with a fixed seed, so every reported analysis is reproducible from the delivered reads. "
        "Variant-call data underlying all figures are provided in the Supplementary Tables. SHMtrace (v0.2) and all "
        "analysis and figure-generation code will be made publicly available upon publication."),
}

# ---------------------------------------------------------------- targeted edits (start, old, new)
EDITS = [
    ("2.7 Read pre-processing", "Read pre-processing", "Read preparation"),
    ("Quality-control-passing reads were analysed for each sample",
     "Quality-control-passing reads were analysed for each sample with SHMtrace v0.2, written for this study "
     "(Data and code availability). Reads were aligned with minimap2 (map-ont) to the 1,230-nt JH4 analysis "
     "reference or the 655-nt N114 analysis reference, which were distinct from the shorter mapping references "
     "used during preprocessing.",
     "Variants were called in the same SHMtrace run, from the 500 usable reads drawn per library and their "
     "alignments to the 1,230-nt JH4 or 655-nt N114 reference."),
    ("Each candidate was tested against the sequencing background",
     "at face value, 56 of 727 such JH4 tests and 72 of 580 N114 tests reached p < 0.01",
     "at face value, 55 of 1,391 such JH4 tests and 69 of 1,225 N114 tests reached p < 0.01"),
    ("Each candidate was tested against the sequencing background",
     "The same reads were originally analysed in Geneious Prime 2026.1.2 (Kearse et al., 2012) with an unadjusted "
     "approximate per-candidate p ≤ 0.05; those calls are retained in the data deposit and are compared in Section "
     "3.1.",
     "An earlier Geneious analysis of separately preprocessed subsets is retained in the data deposit and compared "
     "in Section 3.1."),
    ("Germline and strain-background positions were handled",
     "SHMtrace additionally removed reads carrying the recurrent-contaminant allele at ≥ 50% of its diagnostic "
     "sites. Competitive allele selection and the recurrent-contaminant screens reduce,",
     "Competitive alignment against the paralogue and contaminant references and the paralogue-allele filter "
     "(Section 2.7) reduce,"),
    ("For N114, the 651-nt reference used for preprocessing",
     "For N114, the 651-nt reference used for preprocessing selection and the 655-nt parental sequence used for "
     "variant analysis were not identical records. Variant positions and denominators in this manuscript refer to "
     "the 655-nt analysis reference.",
     "For N114, reads were selected and analysed against the same 655-nt parental reference, to which all variant "
     "positions and denominators refer."),
    ("For the separate full-length ONT analysis",
     "(leave-one-control-out factor 10)", "(leave-one-control-out factor 5)"),
    ("Depth-standardised unique-variant densities were log10-transformed",
     "(one Aicda−/− ONT library and two empty-vector subclones)",
     "(one Aicda−/− ONT library and all three empty-vector subclones)"),
    ("To test whether ONT could recover the expected group-level SHM signal",
     "FASTQ reads were processed through the assay-specific pipeline",
     "The delivered FASTQ reads were analysed end to end by SHMtrace"),
    ("We selected 500 reads as a pragmatic matched working depth",
     "and the deposited subsampling realisation", "and a single random draw per depth"),
    ("This study provides proof-of-concept evidence",
     "assay-specific reference selection, amplicon-length and expected-error filters, matched depth, per-candidate "
     "filtering,",
     "competitive alignment against paralogue references, read-length and expected-error gates, matched depth, "
     "per-candidate testing against the controls,"),
    ("The present data do not establish the accuracy",
     ", and it cannot be calibrated at very low depth (Supplementary Figure S6)", ""),
    ("The workflow also scales favourably",
     "(a mean of 1,586 reads per sample), the corresponding value is ~US$0.022 per read, an approximately 385-fold "
     "difference. Thus, the 385-fold estimate",
     "(a mean of 1,620 usable reads per sample), the corresponding value is ~US$0.022 per read, an approximately "
     "394-fold difference. Thus, the 394-fold estimate"),
    ("Table 1. Illustrative per-sample comparison",
     "uses the mean 1,586 reads available after pre-processing", "uses the mean of 1,620 usable reads per library"),
    ("Supplementary Table S1. Per-sample sequencing",
     "Reads delivered by the sequencing provider and reads satisfying assay-specific target selection, "
     "amplicon-length, and expected-error criteria are shown for ONT libraries;",
     "Reads delivered by the sequencing provider and the usable reads left after SHMtrace's read preparation "
     "(length and expected-error gates, competitive alignment, removal of chimeric and paralogue reads) are shown "
     "for ONT libraries, of which 500 per library were drawn at random for analysis;"),
    ("Group geometric means with descriptive 95% percentile-bootstrap intervals",
     "Wt ONT 59 [26–134]; Aicda−/− Sanger 23 [17–38]; Aicda−/− ONT 0.55 [0.38–1.14] (one library with no calls, "
     "at its one-event density); empty vector 0.33 (two subclones with no calls, at their one-event density, and "
     "one with one call); pMX-Aicda 29 [22–50].",
     "Wt ONT 66 [28–132]; Aicda−/− Sanger 23 [17–38]; Aicda−/− ONT 0.48 [0.38–0.76] (one library with no calls, "
     "at its one-event density); empty vector 0.33 (no subclone had a call; all three at their one-event density); "
     "pMX-Aicda 30 [20–55]."),
    # ---- legends
    ("Figure 1. ",
     "(3) Analysis with SHMtrace. Reads from a related gene copy, joined (chimeric) reads and low-quality reads were "
     "removed; a change was kept only if it appeared more often than sequencing errors do in the AID-null controls, "
     "with at most 5% of kept calls expected to be errors; and groups were compared with one value per mouse or "
     "subclone. Sanger clones took the same steps, keeping only confident base calls.",
     "(3) Analysis with SHMtrace, from the delivered reads. Over-length, low-quality, chimeric and paralogous reads "
     "were removed and 500 usable reads drawn per library; each candidate substitution was tested against the error "
     "rate measured in the AID-null controls (Benjamini–Hochberg false-discovery rate ≤ 0.05); and groups were "
     "compared with one value per mouse or subclone. Sanger clones were analysed in SHMtrace's Sanger mode, counting "
     "substitutions at base quality ≥ 20."),
    ("Figure 2. Depth-standardised", "gave ≥ 108-fold (ONT; p = 0.0018)", "gave ≥ 137-fold (ONT; p = 0.0025)"),
    ("Figure 2. Depth-standardised", "Wt ONT 59 [26–134], Aicda−/− ONT 0.55 [0.38–1.14]",
     "Wt ONT 66 [28–132], Aicda−/− ONT 0.48 [0.38–0.76]"),
    ("Figure 2. Depth-standardised", "ONT libraries were analysed at 500 reads per sample",
     "ONT libraries were analysed at 500 reads per sample, drawn at random from each library's usable reads"),
    ("Figure 3. ONT yields", "n = 471 Sanger and 571 ONT", "n = 471 Sanger and 618 ONT"),
    ("Figure 3. ONT yields", "interaction was p = 0.50", "interaction was p = 0.61"),
    ("Figure 3. ONT yields", "(Pearson r = 0.91; slope = 1.02;", "(Pearson r = 0.93; slope = 1.00;"),
    ("Figure 3. ONT yields", "adjusted paired p = 0.10 and 0.97", "adjusted paired p = 0.15 and 0.98"),
    ("Figure 4. Filtered apparent variants",
     "Two empty-vector subclones had no calls and the third had one; those without calls are drawn (open circles) "
     "and tested at their one-event density (0.33 × 10−5), so the group difference is a lower bound: the "
     "exploratory Welch test on log10-transformed densities gave p = 0.0034 and at least an 89-fold difference "
     "(29 versus 0.33 × 10−5).",
     "None of the three empty-vector subclones had a call; they are drawn (open circles) and tested at their "
     "one-event density (0.33 × 10−5), so the group difference is a lower bound: the exploratory Welch test on "
     "log10-transformed densities gave p = 0.0048 and at least a 92-fold difference (30 versus 0.33 × 10−5)."),
    ("Figure 4. Filtered apparent variants", "(n = 287) and empty vector (n = 1)",
     "(n = 300) and empty vector (n = 0)"),
    ("Supplementary Figure S1. Pooled substitution", "(n = 571 distinct calls)", "(n = 618 distinct calls)"),
    ("Supplementary Figure S1. Pooled substitution", "1,141 bp; n = 680", "1,141 bp; n = 795"),
    ("Supplementary Figure S1. Pooled substitution", "(Pearson r = 0.974; all differences ≤ 2.2 percentage points)",
     "(Pearson r = 0.986; all differences ≤ 1.8 percentage points)"),
    ("Supplementary Figure S3. Individual JH4", "(n = 155, 347, and 69 unique variants, respectively)",
     "(n = 205, 340, and 73 unique variants, respectively)"),
    ("Supplementary Figure S3. Individual JH4", "Aicda−/− mice 1–3 (n = 0, 3, and 1)", "Aicda−/− mice 1–3 (n = 0, 2, and 1)"),
    ("Supplementary Figure S3. Individual JH4", "analysed at ~500 reads per mouse",
     "analysed at 500 reads per mouse, drawn at random from each library's usable reads"),
    ("Supplementary Figure S4. Individual N114", "(n = 69, 68, and 150 unique variants, respectively)",
     "(n = 75, 59, and 166 unique variants, respectively)"),
    ("Supplementary Figure S4. Individual N114", "C3, C4, and D1 (n = 0, 0, and 1)", "C3, C4, and D1 (n = 0, 0, and 0)"),
    ("Supplementary Figure S4. Individual N114", "analysed at ~500 reads per subclone",
     "analysed at 500 reads per subclone, drawn at random from each library's usable reads"),
    ("Supplementary Figure S5. Recurrence-weighted", "(n = 1,455 read-weighted observations); (B) three empty-vector "
     "controls (n = 4)", "(n = 1,478 read-weighted observations); (B) three empty-vector controls (n = 0; no calls)"),
    ("Supplementary Figure S9. Sanger and ONT agree", "(Pearson r = 0.88)", "(Pearson r = 0.92)"),
    ("Supplementary Figure S10. Mutation tree", "(84 reads sharing a V(D)J junction;", "(69 reads sharing a V(D)J junction;"),
    ("Supplementary Figure S10. Mutation tree",
     "87% of the clone's mutated reads fit the tree, and 10 of 1,128 mutation pairs fail the four-gamete test.",
     "70% of the clone's mutated reads fit the tree, and 7 of 1,176 mutation pairs fail the four-gamete test."),
]

# legend bodies replaced whole (the bold label run is kept)
LEGENDS = {
    "Supplementary Figure S6. ": (
        "Sequencing depth versus filtered position–substitution discovery and depth-standardised unique-variant "
        "density. JH4 and N114 libraries were analysed at 30, 60, 125, 250, 500 and 1,000 reads, SHMtrace drawing one "
        "random subset of each library's usable reads at each depth (fixed seed; Wt4, with 962 usable reads, "
        "contributes all of them at 1,000); variants were re-called at each depth with the controls calibrated at the "
        "same depth. (A, C) Distinct calls versus depth; open circles, geometric mean; whiskers, geometric SD (n = 3 "
        "per group). AID-active samples rose toward the 1,000-read count; JH4 controls retained 0–3 calls at every "
        "depth, and N114 controls 0–5 calls at 30 and 60 reads and none from 125 reads upward. (B, D) Distinct calls "
        "divided by depth and analysed bases (× 10−5; geometric mean ± geometric SD; zeros set to one only for the "
        "geometric calculation). The metric declines with depth because its numerator saturates. Dashed line, "
        "500-read working depth. One realised subset was analysed at each sample/depth combination; these "
        "operational curves do not estimate sensitivity against a known truth set."),
    "Supplementary Figure S7. ": (
        "Filtered apparent-variant discovery approached the 1,000-read count by 500 reads in this dataset. For each "
        "AID-active sample, the number of distinct position–substitution calls at each depth is expressed relative "
        "to the maximum observed at any analysed depth. (A) JH4, Wt mice; (B) N114, pMX-Aicda subclones (n = 3; "
        "mean ± SD). At 500 reads, mean counts were 94% (JH4) and 91% (N114) of each sample's maximum. Controls were "
        "omitted from normalisation because their low counts make percentages unstable. One random draw of usable "
        "reads was made per library and depth (fixed seed); no repeated-draw uncertainty analysis was performed. The "
        "curves describe repeatedly detectable filtered positions in these finite libraries. They do not estimate "
        "sensitivity to the unknown true mutation set, establish a universal saturation depth, or show that residual "
        "control calls are true mutations. Dashed line, 500 reads; dotted line, 100%. Absolute discovery is shown in "
        "Supplementary Figure S6."),
    "Supplementary Figure S11. ": (
        "Detailed workflow, as run by SHMtrace from the delivered reads. (A) Experimental design. JH4: Peyer's patch "
        "B220+GL7+Fas+ germinal-centre B cells were sorted from 4-month-old Wt and Aicda−/− mice (n = 3 per "
        "genotype), and the 1,230-bp JH4 intron amplicon was amplified with SuperFi II for ONT and, in separate PCRs, "
        "cloned for Sanger sequencing. ONT analysis covers positions 90–1,230 (1,141 bp); the window shared with "
        "Sanger is positions 101–629 (529 bp). N114: cells were transduced with pMX-Aicda or the empty vector (both "
        "GFP and puromycin resistance), sorted for GFP-high expression, subcloned by limiting dilution and kept under "
        "puromycin for 1 month; three subclones per vector were sequenced, and positions 21–635 of the 655-bp "
        "reference (615 bp) were analysed. Amplicons are drawn to scale. (B) Read preparation in SHMtrace: (1) reads "
        "longer than 2,000 bp (JH4) or 900 bp (N114) were removed, with no minimum length, as were reads averaging "
        "more than 0.01 expected errors per base; (2) the remaining reads were aligned with minimap2 to the target "
        "together with competing references (for JH4, a JH4-like paralogue and a recurrent contaminant haplotype), "
        "and reads aligning best to a competing reference, chimeric reads and reads aligned over less than 800 bp "
        "(JH4) or 400 bp (N114) were removed; (3) JH4 reads carrying ten or more paralogue-specific bases were "
        "removed; (4) 500 of the 962–2,905 usable reads per library were drawn at random (fixed seed). (C) Variant "
        "calling in SHMtrace: (5) every base differing from the reference in the analysed window was a candidate, one "
        "supporting read sufficing; (6) the error rate for each substitution, strand and base quality was measured in "
        "the AID-null controls and multiplied by a five-fold margin chosen by leave-one-control-out; (7) candidates "
        "supported by more reads than these errors explain were kept at a Benjamini–Hochberg false-discovery rate of "
        "0.05, and sites at which ≥ 50% of reads differ from the reference were removed as germline or parental; "
        "(8) results were summarised per animal (distinct calls per base sequenced, spectra, clones and lineage "
        "trees), each mouse or subclone counting once. (D) Sanger clones in SHMtrace's Sanger mode: 447 chromatograms "
        "from eight sequencing orders were assigned to mice with a sample sheet and quality-control flags, trimmed to "
        "their clean core, aligned to JH4, and substitutions called at base quality ≥ 20. (E) The fate of every "
        "delivered read: usable reads (cases blue, AID-null controls red), reads removed by the length and quality "
        "gates (dark grey), and reads removed at alignment or by the paralogue-allele filter (light grey); the line "
        "marks the 500 reads analysed, and numbers give usable reads. Drawn in SciCanvas; no generative image model "
        "was used."),
}

IMAGES = [
    ("Figure 1. ", S1F / "Figure1_overview_preview.png"),
    ("Figure 2. Depth-standardised", FIG / "Figure2_JH4_combined.png"),
    ("Figure 3. ONT yields", FIG / "Figure3_concordance.png"),
    ("Figure 4. Filtered apparent variants", FIG / "Figure4_hybridoma_combined.png"),
    ("Supplementary Figure S1. Pooled substitution", FIG / "FigureS_spectrum_529_vs_fulllength.png"),
    ("Supplementary Figure S3. Individual JH4", FIG / "FigureS6_individual_JH4_ONT_matrices.png"),
    ("Supplementary Figure S4. Individual N114", FIG / "FigureS7_individual_N114_matrices.png"),
    ("Supplementary Figure S5. Recurrence-weighted", FIG / "FigureS7_recurrence_weighted.png"),
    ("Supplementary Figure S6. ", FIG / "FigureS3_depth_discovery_density.png"),
    ("Supplementary Figure S7. ", FIG / "FigureS4_saturation.png"),
    ("Supplementary Figure S9. Sanger and ONT agree", SV / "fig3_agreement_position_hotspot.png"),
    ("Supplementary Figure S10. Mutation tree", CAL / "mutation_trees/mutation_trees_wt3.png"),
    ("Supplementary Figure S11. ", S1F / "Figure_workflow_detailed_preview.png"),
]

S1_HEADER = {4: ("Reads passing filter", "Usable reads")}
TABLES = [  # (table index, old cell text, new)
    (0, "500 reads (down-sampled)", "500 reads (random draw of usable reads)"),
    (0, "observed post-filter Aicda−/− background ~0.55 × 10−5", "observed post-filter Aicda−/− background ~0.48 × 10−5"),
    (2, "p = 0.0018 (≥ 108-fold; the control with no calls at one event)",
     "p = 0.0025 (≥ 137-fold; the control with no calls at one event)"),
    (2, "p = 0.0034 (≥ 89-fold; controls with no calls at one event)",
     "p = 0.0048 (≥ 92-fold; all three controls without calls, at one event)"),
    (2, "r = 0.91 [0.69–0.97]; slope = 1.02", "r = 0.93 [0.77–0.98]; slope = 1.00"),
    (2, "interaction p = 0.50; none significant (minimum adjusted p = 0.46)",
     "interaction p = 0.61; none significant (minimum adjusted p = 0.68)"),
    (2, "Ts adjusted p = 0.10; Tv adjusted p = 0.97", "Ts adjusted p = 0.15; Tv adjusted p = 0.98"),
    (2, "≈0.55 × 10−5 per analysed base", "≈0.48 × 10−5 per analysed base"),
]
# Supplementary Table S1 ONT rows: table key -> (run, library)
S1_ROWS = {"Wt2": ("jh4", "WT2"), "Wt3": ("jh4", "WT3"), "Wt4": ("jh4", "WT4"), "KO1": ("jh4", "KO1"),
           "KO2": ("jh4", "KO2"), "KO3": ("jh4", "KO3"), "1_C3": ("n114", "C3"), "3_C4": ("n114", "C4"),
           "4_D1": ("n114", "D1"), "6_C1": ("n114", "C1"), "7_D3": ("n114", "D3"), "10_D5": ("n114", "D5")}


def s1_values():
    out = {}
    for key, (run, lib) in S1_ROWS.items():
        acct = json.load(open(CAL / f"{run}_shmtrace/table_s1.json"))[lib]
        proj = {s["sample"]: s for s in json.load(open(CAL / f"{run}_shmtrace/project.json"))["samples"]}[lib]
        d = proj["unique_density_x1e5"]
        dens = f"{d:.1f}" if d >= 10 else f"{d:.2f}"
        out[key] = (f"{acct['delivered']:,}", f"{acct['usable']:,}", f"{acct['analysed']}",
                    f"{acct['candidates_tested']}", f"{acct['calls']}", dens)
    return out


def main(src, out):
    doc = Document(src)

    def find(start):
        hits = [p for p in doc.paragraphs if p.text.startswith(start)]
        if len(hits) != 1:
            raise RuntimeError(f"{len(hits)} paragraphs start with {start!r}")
        return hits[0]

    # find every target first, so a rewritten paragraph cannot shadow a later anchor
    whole = [(find(s), new) for s, new in PARAS.items()]
    edits = [(find(s), old, new) for s, old, new in EDITS]
    legends = [(find(s), s, new) for s, new in LEGENDS.items()]
    for p, new in whole:
        body = "".join(r.text for r in p.runs)
        if body != p.text:
            raise RuntimeError(f"paragraph has non-run content: {p.text[:60]!r}")
        replace_text(p, body, new)
    for p, old, new in edits:
        replace_text(p, old, new)
    for p, label, new in legends:
        body = "".join(r.text for r in p.runs)
        replace_text(p, body[len(label):], new)

    for ti, old, new in TABLES:
        cells = [c for row in doc.tables[ti].rows for c in row.cells if old in c.text]
        if not cells:
            raise RuntimeError(f"table {ti}: not found {old!r}")
        set_cell(cells[0], old, new)
    s1 = doc.tables[1]
    for col, (old, new) in S1_HEADER.items():
        set_cell(s1.rows[0].cells[col], old, new)
    vals, seen = s1_values(), set()
    for row in s1.rows:
        key = row.cells[0].text.strip()
        if key in vals and row.cells[2].text.strip() == "ONT":
            seen.add(key)
            for col, val in zip((3, 4, 5, 7, 8, 9), vals[key], strict=True):
                cell = row.cells[col]
                if cell.text.strip() != val:
                    set_cell(cell, cell.text.strip(), val)
    if seen != set(vals):
        raise RuntimeError(f"S1 rows missing: {set(vals) - seen}")

    for start, path in IMAGES:
        legend = find(start)
        idx = [p._p for p in doc.paragraphs].index(legend._p)
        img_p = next(doc.paragraphs[k] for k in range(idx - 1, idx - 4, -1)
                     if list(doc.paragraphs[k]._p.iter(qn("a:blip"))))
        swap_image(doc, img_p, path)

    for p in doc.paragraphs:
        restyle(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    restyle(p)
    doc.core_properties.comments = ("Ver21: all ONT results from SHMtrace run end to end on the delivered reads "
                                    "(shmtrace docs/calibration/manuscript_corrected/revise_manuscript_v21.py)")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
