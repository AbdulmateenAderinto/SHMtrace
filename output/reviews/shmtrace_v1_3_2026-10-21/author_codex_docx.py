"""Author the prepared manuscript with bundled python-docx, without analysis dependencies."""
import json
import sys
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Mm, Pt, RGBColor

HERE=Path(__file__).parent.resolve()
DEST=HERE/'codex_version'
sys.path.insert(0,str(DEST/'sources'))
from docbuild import MANUAL, Manuscript  # noqa: E402

MANUAL['seija2026'] = (MANUAL['seija2026'].replace('Häfner, K.,','Häfner, K.A.,')
    .replace('Gemeinhardt, T.,','Gemeinhardt, T.M.,').replace('Subramani, P.,','Subramani, P.G.,')
    .replace('Piskor, E.,','Piskor, E.-M.,'))


class CodexManuscript(Manuscript):
    def __init__(self):
        super().__init__(line_numbers=False,links=True,footer_note='SHMtrace · v1.3 Codex 1')
        section=self.doc.sections[0]
        section.left_margin=section.right_margin=Mm(20.95)
        self.doc.core_properties.title='SHMtrace Codex revision'
        self.doc.core_properties.subject='Complete manuscript with supplemental information'
        for style in self.doc.styles:
            if style.name.startswith(('Title','Heading')):
                style.font.color.rgb=RGBColor(0,0,0)

    def figure(self,path,width_mm=174):
        paragraph=super().figure(path,width_mm=174)
        paragraph.paragraph_format.keep_with_next=True
        paragraph.paragraph_format.space_after=Pt(5)
        return paragraph

    def caption(self,title,legend,size=10):
        # Images are kept at their intended width; a long caption may continue.
        result=super().caption(title,legend,size=size)
        return result

    def figure_block(self,paths,title,legend,max_width_mm=174,min_width_mm=174):
        from PIL import Image
        # The two-image supplementary entries may span pages. Preserve label size
        # rather than reducing both panels and their text to fit a single page.
        for i,path in enumerate(paths):
            im=Image.open(path)
            height=174*im.height/im.width
            width=174 if height<=216 else 174*216/height
            paragraph=super().figure(path,width_mm=width)
            paragraph.paragraph_format.keep_with_next=(i==len(paths)-1)
            if i<len(paths)-1:
                self.page_break()
        self.caption(title,legend)
        return 174

    def table(self,header,rows,widths=None,size=8.5,title=None,footnote=None):
        if len(header)==15 and header[:2]==['Assay','Library']:
            # Fifteen columns made single words and numerical values break across
            # lines. Split the same records into two aligned tables without loss.
            self.table(['Assay','Library','Delivered','Over-length','High error',
                'Unmapped / paralog','Chimeric / short','Paralog alleles','Usable','Analyzed'],
                [row[:10] for row in rows],
                widths=[12,14,18,16,16,24,22,20,16,16],size=size,title=title)
            self.para('Calling results',size=9.5,bold=True,space_after=4)
            self.table(['Assay','Library','Candidates tested','Calls (FDR ≤ 0.05)',
                'Unique density (×10^{−5})','Read-weighted frequency (×10^{−3})',
                'Clone-weighted frequency (×10^{−3})'],
                [[row[i] for i in (0,1,10,11,12,13,14)] for row in rows],
                widths=[15,16,28,25,30,30,30],size=size,footnote=footnote)
            return
        if title:
            title=title.replace('Figures S14 and S9','Figures S14 and S15')
        if footnote:
            footnote=footnote.replace('analyses of Figures 2 and 4','analyses of Figures 4 and 5')
        # Proportionally fit each pre-existing table to the 174-mm text area.
        if widths:
            factor=174/sum(widths)
            widths=[w*factor for w in widths]
        super().table(header,rows,widths=widths,size=size,title=title,footnote=footnote)
        table=self.doc.tables[-1]
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False

    def references(self,heading='References'):
        before=len(self.doc.paragraphs)
        super().references(heading)
        for paragraph in self.doc.paragraphs[before+1:]:
            paragraph.paragraph_format.keep_together=True


ops=json.loads((DEST/'document_operations.json').read_text())
manuscript=CodexManuscript()
main_figures=0
for name,args,kwargs in ops:
    if name=='figure':
        if main_figures:
            manuscript.page_break()
        main_figures+=1
    getattr(manuscript,name)(*args,**kwargs)
manuscript.doc.paragraphs[0].style='Title'
target=DEST/'SHMtrace_CellReportsMethods_complete_v1.3_Codex_1.docx'
manuscript.save(target)
print(target)
print('References:',len(manuscript.order),'Main figures:',main_figures)
