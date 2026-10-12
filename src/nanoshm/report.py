"""Exports for a project result: CSV, Excel, figures, a standalone HTML report, ZIP.

Everything here takes the dict `project.run_project` returns.
"""
from __future__ import annotations

import csv
import io
import json
import os
import zipfile
from datetime import date

from . import viz

SAMPLE_COLUMNS = ["sample", "group", "control", "reads", "unique_mutations",
                  "unique_density_x1e5", "one_event_density_x1e5",
                  "combined_density_x1e5", "combined_density_ci_low",
                  "combined_density_ci_high", "multi_read_mutations",
                  "single_read_observed", "single_read_background", "single_read_estimate",
                  "single_read_density_x1e5", "single_read_density_ci_low",
                  "single_read_density_ci_high",
                  "indels_called", "insertions", "deletions", "frameshift_indels",
                  "indel_density_x1e5",
                  "mutation_frequency_x1e3", "mutated_reads_pct",
                  "mean_mutations_per_read", "pct_GC", "pct_transitions",
                  "pct_GC_in_WRC_GYW", "pct_AT_in_WA_TW", "analysed_bp", "null_from",
                  "model_used", "warnings"]
VARIANT_COLUMNS = ["sample", "group", "pos", "ref", "alt", "reads", "depth", "frequency",
                   "p_value", "plus_reads", "minus_reads", "context", "homopolymer",
                   "WRC_GYW"]


def _rows_csv(rows: list[dict], columns: list[str] | None = None) -> str:
    if not rows:
        return ""
    columns = columns or list(rows[0].keys())
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def tables(res: dict) -> dict[str, str]:
    """Every table as CSV text, keyed by file name."""
    spec_rows = []
    for g, c in res["spectrum_group"].items():
        tot = sum(c.values()) or 1
        for k, v in sorted(c.items()):
            spec_rows.append({"group": g, "substitution": k, "count": v,
                              "percent": round(100 * v / tot, 2)})
    load_rows = []
    for s, hist in res["mutation_load"].items():
        for k in sorted(hist, key=int):
            load_rows.append({"sample": s, "mutations_per_read": k, "reads": hist[k]})
    return {
        "samples.csv": _rows_csv(res["samples"], SAMPLE_COLUMNS),
        "groups.csv": _rows_csv(res["groups"]),
        "variants.csv": _rows_csv(res["variants"], VARIANT_COLUMNS),
        **({"indels.csv": _rows_csv(res["indels"])} if res.get("indels") else {}),
        "spectrum.csv": _rows_csv(spec_rows),
        "positional_frequency.csv": _rows_csv(res["positional"]),
        "mutation_load.csv": _rows_csv(load_rows),
    }


def workbook(res: dict) -> bytes:
    """All tables as one .xlsx (needs openpyxl)."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {"Samples": (res["samples"], SAMPLE_COLUMNS),
              "Groups": (res["groups"], None),
              "Variants": (res["variants"], VARIANT_COLUMNS),
              "Positional": (res["positional"], None)}
    if res.get("indels"):
        sheets["Indels"] = (res["indels"], None)
    spec_rows = [{"group": g, "substitution": k, "count": v}
                 for g, c in res["spectrum_group"].items() for k, v in sorted(c.items())]
    sheets["Spectrum"] = (spec_rows, None)
    for title, (rows, cols) in sheets.items():
        ws = wb.create_sheet(title)
        if not rows:
            continue
        cols = cols or list(rows[0].keys())
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = \
                min(28, max(10, len(str(c)) + 2))
    meta = wb.create_sheet("Run", 0)
    meta.append(["nanoshm report", res.get("name", "")])
    meta.append(["date", str(date.today())])
    for k, v in res["params"].items():
        meta.append([k, str(v)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def figures(res: dict, fmt: str = "png") -> dict[str, bytes]:
    """Publication figures for a project, keyed by file name."""
    return viz.project_figures(res, fmt=fmt)


def _data_uri(b: bytes, fmt: str) -> str:
    import base64
    mime = {"png": "image/png", "svg": "image/svg+xml", "pdf": "application/pdf"}[fmt]
    return f"data:{mime};base64," + base64.b64encode(b).decode()


def _html_table(rows: list[dict], columns: list[str] | None = None, limit: int = 2000) -> str:
    if not rows:
        return "<p class=muted>Nothing to show.</p>"
    columns = [c for c in (columns or rows[0].keys()) if any(c in r for r in rows)]
    head = "".join(f"<th>{c.replace('_', ' ')}</th>" for c in columns)
    body = []
    for r in rows[:limit]:
        tds = []
        for c in columns:
            v = r.get(c)
            v = "" if v is None else ("yes" if v is True else "no" if v is False else v)
            tds.append(f"<td>{v}</td>")
        body.append("<tr>" + "".join(tds) + "</tr>")
    more = (f"<p class=muted>Showing the first {limit} of {len(rows)} rows; the CSV export "
            "has them all.</p>") if len(rows) > limit else ""
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>{more}"


def html_report(res: dict, fmt: str = "png") -> str:
    """One self-contained HTML file: figures, tables and the run's settings."""
    figs = figures(res, fmt=fmt)
    fig_html = "".join(
        f"<figure><img src='{_data_uri(v, fmt)}' alt='{k}'><figcaption>{k.rsplit('.', 1)[0].replace('_', ' ')}"
        f"</figcaption></figure>" for k, v in figs.items())
    p = res["params"]
    notes_html = "".join(
        f"<p class=muted><b>Note:</b> {n}</p>" for n in res.get("notes") or [])
    settings = "".join(f"<tr><th>{k.replace('_', ' ')}</th><td>{v}</td></tr>" for k, v in p.items())
    return f"""<!doctype html>
<html lang=en><head><meta charset=utf-8>
<title>{res.get('name', 'nanoshm report')}</title>
<style>
 :root {{ color-scheme: light dark; --fg:#111; --muted:#666; --line:#ddd; --bg:#fff; }}
 @media (prefers-color-scheme: dark) {{ :root {{ --fg:#eee; --muted:#999; --line:#333; --bg:#111; }} }}
 body {{ font: 15px/1.5 -apple-system, system-ui, sans-serif; color:var(--fg);
         background:var(--bg); margin:0 auto; padding:24px 16px; max-width:1100px; }}
 h1 {{ font-size:22px; margin:0 0 4px; }} h2 {{ font-size:17px; margin:28px 0 8px; }}
 .muted {{ color:var(--muted); }}
 table {{ border-collapse:collapse; width:100%; font-size:13px; margin:6px 0 14px; }}
 th,td {{ border-bottom:1px solid var(--line); padding:4px 8px; text-align:left;
          white-space:nowrap; }}
 thead th {{ position:sticky; top:0; background:var(--bg); }}
 figure {{ margin:0 0 18px; }} img {{ max-width:100%; height:auto; }}
 figcaption {{ color:var(--muted); font-size:13px; }}
 .wrap {{ overflow-x:auto; }}
</style></head><body>
<h1>{res.get('name', 'nanoshm report')}</h1>
<p class=muted>{date.today()} &middot; {p.get('region')} &middot; {p.get('platform')} &middot;
 mode {p.get('mode')} &middot; {len(res['samples'])} samples</p>
{notes_html}
<h2>Figures</h2>{fig_html}
<h2>Samples</h2><div class=wrap>{_html_table(res['samples'], SAMPLE_COLUMNS)}</div>
<h2>Groups</h2><div class=wrap>{_html_table(res['groups'])}</div>
<h2>Mutations</h2><div class=wrap>{_html_table(res['variants'], VARIANT_COLUMNS)}</div>
<h2>Settings</h2><table>{settings}</table>
</body></html>"""


def write_all(res: dict, out_dir: str, fmt: str = "png", excel: bool = True) -> list[str]:
    os.makedirs(out_dir, exist_ok=True)
    written = []
    for name, text in tables(res).items():
        path = os.path.join(out_dir, name)
        with open(path, "w", newline="") as fh:
            fh.write(text)
        written.append(path)
    for name, blob in figures(res, fmt=fmt).items():
        path = os.path.join(out_dir, name)
        with open(path, "wb") as fh:
            fh.write(blob)
        written.append(path)
    path = os.path.join(out_dir, "report.html")
    with open(path, "w") as fh:
        fh.write(html_report(res, fmt="png"))
    written.append(path)
    path = os.path.join(out_dir, "project.json")
    with open(path, "w") as fh:
        json.dump({k: v for k, v in res.items() if k not in ("reference", "_details")},
                  fh, indent=2, default=str)
    written.append(path)
    if excel:
        try:
            path = os.path.join(out_dir, "tables.xlsx")
            with open(path, "wb") as fh:
                fh.write(workbook(res))
            written.append(path)
        except ImportError:
            pass
    return written


def zip_bytes(res: dict, fmt: str = "png") -> bytes:
    """Everything in one zip, for the GUI's download button."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in tables(res).items():
            z.writestr(name, text)
        for name, blob in figures(res, fmt=fmt).items():
            z.writestr(name, blob)
        z.writestr("report.html", html_report(res, fmt="png"))
        z.writestr("project.json", json.dumps(
            {k: v for k, v in res.items() if k not in ("reference", "_details")},
            indent=2, default=str))
        try:
            z.writestr("tables.xlsx", workbook(res))
        except ImportError:
            pass
    return buf.getvalue()
