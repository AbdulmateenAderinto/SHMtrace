"""A sample sheet: which file belongs to which sample, written down once.

Folders and file names are a fragile record of which clone came from which
animal: a clone filed in the wrong folder, or a trace copied under another
name, silently changes a sample's result. A sample sheet makes the assignment
explicit and checkable. It is a CSV (or tab-separated) table, one row per file:

    file,sample,group,control,order,well,label,exclude
    traces/434141_A7_w2-1.ab1,WT2,WT,no,434141,A7,w2-1,
    traces/434141_E8_w2-1.ab1,WT2,WT,no,434141,E8,w2-1,
    traces/432625_E6_a1-1.ab1,KO1,KO,yes,432625,E6,a1-1,AID-type mutations in a control

Columns (names are case-insensitive; spaces and dashes count as underscores):

  file      required. A read file or a folder of them, relative to the sheet's
            own folder or absolute. A pattern with * or ? matches several.
  sample    required. The animal or independent source the file belongs to.
  group     the sample's group (default: the sample's own name).
  control   yes/no: a mutation-free control (e.g. an AID-knockout).
  clone_id  what to call the clone in every table. Default: order-well when
            both are given (labels are reused across sequencing orders, wells
            within an order are not), otherwise the file name. Two traces from
            one order and well (a re-run) become order-well-<file name>.
  order, well, label
            where the trace came from: sequencing order, plate well, and the
            label written on the tube. Kept with every flag and exclusion.
  exclude   any text here leaves the row out of the analysis, with the text
            recorded as the reason.
  barcode, barcode2
            inline barcode(s) of a pooled run: every such row names the same
            pooled FASTQ in `file`, and the run is split by barcode before the
            analysis (demux.py). barcode2 is the barcode at the other end.
  aliases   how this sample is written on tube labels, separated by ; (for
            WT3: w3;wt3). Optional; with it, a clone whose label names a
            different sample is flagged (qc.py, label_mismatch).

Any other column is kept as metadata. `load` checks the sheet and reports
every problem at once (with its line number) instead of stopping at the first.
"""
from __future__ import annotations

import csv
import glob
import hashlib
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from .io_utils import READ_EXTENSIONS, _stem, sniff_format

REQUIRED = ("file", "sample")
KNOWN = ("file", "sample", "group", "control", "clone_id", "order", "well", "label", "exclude",
         "barcode", "barcode2",
         "aliases")
YES = {"yes", "y", "true", "t", "1", "control", "x"}
NO = {"no", "n", "false", "f", "0", ""}


class SheetError(ValueError):
    """The sheet cannot be used; `problems` lists every reason, by line."""

    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("sample sheet problems:\n  " + "\n  ".join(problems))


@dataclass
class Row:
    line: int
    file: str              # resolved, absolute
    sample: str
    group: str
    control: bool
    clone_id: str
    order: str = ""
    well: str = ""
    label: str = ""
    exclude: str = ""
    barcode: str = ""        # inline barcode(s): the file is a pooled run (demux.py)
    barcode2: str = ""
    id_given: bool = True
    extra: dict = field(default_factory=dict)

    def provenance(self) -> dict:
        return {"clone_id": self.clone_id, "sample": self.sample, "file": self.file,
                "order": self.order, "well": self.well, "label": self.label,
                "sheet_line": self.line}


@dataclass
class Sheet:
    path: str
    rows: list[Row]
    warnings: list[str]

    @property
    def used(self) -> list[Row]:
        return [r for r in self.rows if not r.exclude]

    def samples(self) -> list[dict]:
        """The project's sample list: each sample's files, named by clone_id."""
        out: dict[str, dict] = {}
        for r in self.used:
            s = out.setdefault(r.sample, {"name": r.sample, "path": [], "group": r.group,
                                          "control": r.control})
            s["path"].append([r.file, r.clone_id])
        return list(out.values())

    def provenance(self) -> dict[str, dict]:
        """clone_id -> where it came from, for every row (excluded ones too)."""
        return {r.clone_id: r.provenance() for r in self.rows}

    def aliases(self) -> dict[str, list[str]]:
        """sample -> the label tokens that name it (the aliases column, lower case)."""
        out: dict[str, set] = {}
        for r in self.rows:
            toks = {t.strip().lower() for t in r.extra.get("aliases", "").split(";") if t.strip()}
            out.setdefault(r.sample, set()).update(toks)
        return {k: sorted(v) for k, v in out.items() if v}

    def excluded(self) -> list[dict]:
        return [{**r.provenance(), "reason": r.exclude} for r in self.rows if r.exclude]

    def project(self, reference: str, region: str, platform: str = "sanger",
                demux_dir: str | None = None, left_anchor: str | None = None,
                right_anchor: str | None = None, demux_options: dict | None = None,
                **extra) -> dict:
        """A run_project() dict for this sheet.

        Pooled rows (a `barcode` column) are first demultiplexed (demux.py) into one
        FASTQ per sample, in `demux_dir` (default: demux_<pooled file> beside the
        sheet). The barcode sits just outside the primer anchors, by default the
        first and last 20 reference bases (reference-forward).
        """
        samples = self.samples()
        pooled = [r for r in self.used if r.barcode]
        if pooled:
            from collections import defaultdict

            from .demux import demultiplex
            from .io_utils import read_fasta
            ref = read_fasta(reference)[0][1].upper()
            by_file: dict = defaultdict(dict)
            for r in pooled:
                by_file[r.file][r.sample] = (r.barcode, r.barcode2 or None)
            paths, stats = {}, {}
            for f, smp in by_file.items():
                out = demux_dir or os.path.join(os.path.dirname(self.path), f"demux_{_stem(f)}")
                d = demultiplex(f, smp, left_anchor or ref[:20], right_anchor or ref[-20:], out,
                                **(demux_options or {}))
                paths.update(d["paths"])
                stats[os.path.basename(f)] = d["stats"]
            for s in samples:
                if s["name"] in paths:
                    s["path"] = paths[s["name"]]
            extra.setdefault("demux", stats)
        return {"name": Path(self.path).stem, "reference": reference, "region": region,
                "platform": platform, "samples": samples,
                "sample_sheet": {"path": self.path, "provenance": self.provenance(),
                                 "aliases": self.aliases(),
                                 "excluded": self.excluded(), "warnings": self.warnings},
                **extra}


def _norm(h: str) -> str:
    return re.sub(r"[\s\-]+", "_", (h or "").strip().lower())


def _read_table(path: str) -> list[dict]:
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SheetError(["reading .xlsx needs openpyxl; save the sheet as CSV"]) from exc
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            text = fh.read()
        dialect = "excel-tab" if text.count("\t") > text.count(",") else "excel"
        rows = list(csv.reader(text.splitlines(), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise SheetError(["the sheet is empty"])
    head = [_norm(h) for h in rows[0]]
    return [dict(zip(head, [c.strip() for c in r], strict=False)) | {"_line": i + 2}
            for i, r in enumerate(rows[1:])]


def _expand(spec: str, base: Path) -> list[str]:
    p = Path(os.path.expanduser(spec))
    if not p.is_absolute():
        p = base / p
    if any(ch in spec for ch in "*?["):
        return sorted(f for f in glob.glob(str(p)) if os.path.isfile(f))
    if p.is_dir():
        return sorted(str(f) for f in p.iterdir()
                      if f.is_file() and not f.name.startswith(".") and sniff_format(str(f)))
    return [str(p)] if p.exists() else []


def _sha1(path: str) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load(path: str) -> Sheet:
    """Read and check a sample sheet. Raises SheetError listing every problem."""
    path = os.path.abspath(os.path.expanduser(path))
    base = Path(path).parent
    recs = _read_table(path)
    problems: list[str] = []
    warnings: list[str] = []
    cols = set(recs[0]) - {"_line"} if recs else set()
    for c in REQUIRED:
        if c not in cols:
            problems.append(f"no '{c}' column (columns found: {', '.join(sorted(cols)) or 'none'})")
    if problems:
        raise SheetError(problems)

    rows: list[Row] = []
    for rec in recs:
        ln = rec["_line"]
        sample = rec.get("sample", "")
        if not sample:
            problems.append(f"line {ln}: no sample")
            continue
        ctrl_txt = rec.get("control", "").lower()
        if ctrl_txt not in YES | NO:
            problems.append(f"line {ln}: control is '{rec['control']}'; use yes or no")
            continue
        files = _expand(rec.get("file", ""), base) if rec.get("file") else []
        if not files:
            problems.append(f"line {ln}: file not found: {rec.get('file') or '(blank)'}")
            continue
        order, well = rec.get("order", ""), rec.get("well", "").upper()
        for f in files:
            if sniff_format(f) is None:
                problems.append(f"line {ln}: {os.path.basename(f)} is not a read file "
                                f"({', '.join(READ_EXTENSIONS)})")
                continue
            cid = rec.get("clone_id", "")
            if not cid:
                cid = f"{order}-{well}" if order and well else _stem(f)
                if rec.get("barcode"):
                    cid = f"{sample}-{rec['barcode'].upper()}"
                given = False
            else:
                given = True
            if len(files) > 1 and rec.get("clone_id"):
                cid = f"{cid}/{_stem(f)}"
            rows.append(Row(line=ln, file=os.path.abspath(f), sample=sample,
                            group=rec.get("group") or sample, control=ctrl_txt in YES,
                            clone_id=cid, order=order, well=well, label=rec.get("label", ""),
                            exclude=rec.get("exclude", ""), id_given=given,
                            barcode=rec.get("barcode", "").upper(),
                            barcode2=rec.get("barcode2", "").upper(),
                            extra={k: v for k, v in rec.items()
                                   if (k not in KNOWN or k == "aliases") and k != "_line" and v}))

    # default names that collide (a well sequenced twice) take the file name too
    n_default = Counter(r.clone_id for r in rows if not r.id_given)
    for r in rows:
        if not r.id_given and n_default[r.clone_id] > 1:
            r.clone_id = f"{r.clone_id}-{_stem(r.file)}"

    # one file, one row; one name per clone; a well twice in one order is worth a look
    by_file: dict[str, Row] = {}
    by_well: dict[tuple, Row] = {}
    by_id: dict[str, Row] = {}
    pooled_codes: dict[tuple, Row] = {}
    for r in rows:
        if r.barcode:
            # a pooled run: one file, many samples, told apart by barcode
            k = (r.file, r.barcode, r.barcode2)
            if k in pooled_codes:
                problems.append(f"line {r.line}: barcode {r.barcode}{'+' + r.barcode2 if r.barcode2 else ''} "
                                f"is already line {pooled_codes[k].line}'s in this pooled file")
            pooled_codes.setdefault(k, r)
            o = by_file.get(r.file)
            if o is not None and not o.barcode:
                problems.append(f"line {r.line}: {os.path.basename(r.file)} is listed on line "
                                f"{o.line} without a barcode")
            by_file.setdefault(r.file, r)
            continue
        if r.file in by_file:
            o = by_file[r.file]
            problems.append(f"line {r.line}: {os.path.basename(r.file)} is already listed on "
                            f"line {o.line} (sample {o.sample}); a file belongs to one sample")
        by_file.setdefault(r.file, r)
        if r.order and r.well:
            k = (r.order, r.well)
            if k in by_well and by_well[k].file != r.file:
                warnings.append(f"line {r.line}: order {r.order} well {r.well} is also line "
                                f"{by_well[k].line}, a different trace (a re-run, or a well "
                                "recorded wrongly); both are analysed")
            by_well.setdefault(k, r)
        if r.clone_id in by_id and by_id[r.clone_id].file != r.file:
            problems.append(f"line {r.line}: clone_id {r.clone_id} is already used on line "
                            f"{by_id[r.clone_id].line}; give one of them a clone_id")
        by_id.setdefault(r.clone_id, r)

    # a sample has one group and is a control or not
    seen: dict[str, Row] = {}
    for r in rows:
        o = seen.setdefault(r.sample, r)
        if (o.group, o.control) != (r.group, r.control):
            problems.append(f"line {r.line}: sample {r.sample} is group '{r.group}', control "
                            f"{'yes' if r.control else 'no'}, but line {o.line} says group "
                            f"'{o.group}', control {'yes' if o.control else 'no'}")

    # the same trace under two names: a copy, not a second clone
    by_hash: dict[str, Row] = {}
    for r in rows:
        if r.exclude or r.barcode:
            continue
        h = _sha1(r.file)
        if h in by_hash:
            o = by_hash[h]
            where = "the same sample" if o.sample == r.sample else f"sample {o.sample}"
            warnings.append(f"line {r.line} ({r.clone_id}, {r.sample}) is byte-for-byte the same "
                            f"file as line {o.line} ({o.clone_id}, {where}); a copied trace is "
                            "one clone, not two")
        by_hash.setdefault(h, r)

    if problems:
        raise SheetError(problems)
    used = {r.sample for r in rows if not r.exclude}
    for s in sorted({r.sample for r in rows} - used):
        warnings.append(f"every row of sample {s} is excluded, so it is not analysed")
    return Sheet(path=path, rows=rows, warnings=warnings)


_WELL = re.compile(r"(?:^|[_\-\s])([A-H](?:0?[1-9]|1[0-2]))(?=[_\-\s.]|$)")
_ORDER = re.compile(r"(?:^|[_\-\s])(\d{5,})(?=[_\-\s.]|$)")


def draft(folder: str, out: str | None = None) -> str:
    """A first sheet for a folder of read files, to be checked and edited.

    One row per read file. `sample` is the file's sub-folder (blank for files
    at the top), and order and well are guessed from the file name where it
    looks like one. Nothing here is trusted: read every row before using it.
    """
    root = Path(os.path.expanduser(folder)).resolve()
    files = sorted(p for p in root.rglob("*")
                   if p.is_file() and not p.name.startswith(".") and sniff_format(str(p)))
    out = out or str(root / "samplesheet.csv")
    with open(out, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["file", "sample", "group", "control", "order", "well", "label", "exclude"])
        for f in files:
            rel = f.relative_to(root)
            sample = rel.parts[0] if len(rel.parts) > 1 else ""
            stem = _stem(str(f))
            well = _WELL.search(stem)
            order = _ORDER.search(stem)
            w.writerow([str(f), sample, "", "", order.group(1) if order else "",
                        well.group(1) if well else "", stem, ""])
    return out


def main(argv=None):
    import argparse
    import json
    ap = argparse.ArgumentParser(prog="nanoshm-sheet", description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("draft", help="write a first sheet for a folder of read files")
    d.add_argument("folder")
    d.add_argument("-o", "--out")
    c = sub.add_parser("check", help="check a sheet and list its samples")
    c.add_argument("sheet")
    r = sub.add_parser("run", help="analyse every sample in a sheet, with QC flags")
    r.add_argument("sheet")
    r.add_argument("--reference", required=True)
    r.add_argument("--region", required=True, help="preset name, e.g. jh4_intron")
    r.add_argument("--platform", default="sanger")
    r.add_argument("--contaminant", action="append", default=[])
    r.add_argument("--mode", default="recommended")
    r.add_argument("-o", "--out", required=True, help="output folder")
    a = ap.parse_args(argv)
    if a.cmd == "draft":
        print(draft(a.folder, a.out))
        return
    try:
        sh = load(a.sheet)
    except SheetError as e:
        raise SystemExit(str(e)) from None
    for wmsg in sh.warnings:
        print("warning:", wmsg)
    if a.cmd == "check":
        for s in sh.samples():
            print(f"{s['name']:16} group {s['group']:10} {'control' if s['control'] else '':8}"
                  f"{len(s['path'])} file(s)")
        for e in sh.excluded():
            print(f"excluded {e['clone_id']} ({e['sample']}): {e['reason']}")
        return
    from . import project, report
    proj = sh.project(a.reference, a.region, a.platform, contaminants=a.contaminant, mode=a.mode)
    res = project.run_project(proj, lambda f, m: print(f"{100 * f:3.0f}%  {m}"))
    written = report.write_all(res, a.out)
    with open(os.path.join(a.out, "project.json"), "w") as fh:
        json.dump(project.public(res), fh, indent=1, default=str)
    for f in written:
        print(f)


if __name__ == "__main__":
    main()
